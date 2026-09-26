import traceback
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from ortools.sat.python import cp_model
import pandas as pd
import json
import os

def hata_yakalayici(exc_type, exc_value, exc_traceback):
    hata_metni = "".join(traceback.format_exception(exc_type, exc_value, exc_traceback))
    hata_penceresi = tk.Toplevel()
    hata_penceresi.title("Beklenmeyen Sistem Hatası")
    hata_penceresi.geometry("700x400")
    hata_penceresi.attributes('-topmost', True)
    
    tk.Label(hata_penceresi, text="Programda arka planda bir hata oluştu. Lütfen aşağıdaki detayları inceleyin:", 
             font=("Segoe UI", 10, "bold"), fg="#D32F2F").pack(anchor="w", padx=10, pady=10)
    
    frame_text = tk.Frame(hata_penceresi)
    frame_text.pack(expand=True, fill="both", padx=10)
    scrollbar = ttk.Scrollbar(frame_text)
    scrollbar.pack(side="right", fill="y")
    metin_kutusu = tk.Text(frame_text, wrap="word", font=("Consolas", 9), yscrollcommand=scrollbar.set)
    metin_kutusu.insert(tk.END, hata_metni)
    metin_kutusu.config(state="disabled")
    metin_kutusu.pack(side="left", expand=True, fill="both")
    scrollbar.config(command=metin_kutusu.yview)
    
    tk.Button(hata_penceresi, text="Pencereyi Kapat", command=hata_penceresi.destroy, 
              bg="#ececec", font=("Segoe UI", 10)).pack(pady=10)

class DersProgramiApp:
    def __init__(self, root):
        self.root = root
        self.root.title("KMÜ Ders Programı Planlayıcısı")
        self.root.geometry("1150x880")
        
        self.style = ttk.Style()
        self.style.theme_use("clam")
        self.font_standart = ("Segoe UI", 10)
        self.font_baslik = ("Segoe UI", 10, "bold")
        
        self.style.configure("Treeview.Heading", font=self.font_baslik, background="#e1e1e1", foreground="black")
        self.style.configure("Treeview", font=self.font_standart, rowheight=25)
        self.style.configure("TLabel", font=self.font_standart)
        self.style.configure("TLabelframe.Label", font=self.font_baslik, foreground="#333333")
        
        self.veri_dosyasi = "kmu_veritabani.json"
        
        self.bolum_adi = "BÖLÜMÜ"
        self.ogretim_elemanlari = []
        self.siniflar = []
        self.dersler = []
        self.olusturulan_program = []

        self.gunler = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma"]
        self.saatler = [
            "08:30-09:15", "09:30-10:15", "10:30-11:15", "11:30-12:15",
            "13:00-13:45", "14:00-14:45", "15:00-15:45", "16:00-16:45"
        ]
        self.ogle_arasi_siniri = 4

        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=10)

        self.tab_ogr = ttk.Frame(self.notebook)
        self.tab_siniflar = ttk.Frame(self.notebook)
        self.tab_dersler = ttk.Frame(self.notebook)
        self.tab_sonuc = ttk.Frame(self.notebook)

        self.notebook.add(self.tab_ogr, text=" 1. Öğretim Elemanları ")
        self.notebook.add(self.tab_siniflar, text=" 2. Sınıflar & Laboratuvarlar ")
        self.notebook.add(self.tab_dersler, text=" 3. Ders Havuzu ")
        self.notebook.add(self.tab_sonuc, text=" 4. Program Üret & Aktar ")

        self.ogr_musaitlik_vars = {}
        self.sinif_musaitlik_vars = {}
        self.derslik_gerekmez_var = tk.BooleanVar(value=False)
        self.lab_gerekir_var = tk.BooleanVar(value=False)
        self.aktif_mi_var = tk.BooleanVar(value=True)

        self.arayuz_ogretim_elemanlari()
        self.arayuz_siniflar()
        self.arayuz_dersler()
        self.arayuz_sonuc()

        self.veri_yukle()

    def toggle_gun(self, gun, var_dict):
        ilk_durum = var_dict[(gun, self.saatler[0])].get()
        for saat in self.saatler: var_dict[(gun, saat)].set(not ilk_durum)

    def toggle_blok(self, is_sabah, var_dict):
        saatler_dilimi = self.saatler[:self.ogle_arasi_siniri] if is_sabah else self.saatler[self.ogle_arasi_siniri:]
        ilk_durum = var_dict[(self.gunler[0], saatler_dilimi[0])].get()
        for gun in self.gunler:
            for saat in saatler_dilimi: var_dict[(gun, saat)].set(not ilk_durum)

    def tumunu_sec(self, var_dict):
        ilk_durum = var_dict[(self.gunler[0], self.saatler[0])].get()
        for var in var_dict.values(): var.set(not ilk_durum)

    def form_sifirla_ogr(self):
        for item in self.tree_ogr.selection(): self.tree_ogr.selection_remove(item)
        self.entry_ogr_ad.delete(0, tk.END)
        for var in self.ogr_musaitlik_vars.values(): var.set(True)

    def form_sifirla_sinif(self):
        for item in self.tree_sinif.selection(): self.tree_sinif.selection_remove(item)
        self.entry_sinif.delete(0, tk.END)
        self.combo_sinif_tipi.set("Derslik")
        for var in self.sinif_musaitlik_vars.values(): var.set(True)

    def form_sifirla_ders(self):
        for tree in [self.tree_2026, self.tree_2024, self.tree_diger]:
            for item in tree.selection(): tree.selection_remove(item)
        self.entry_ders.delete(0, tk.END)
        self.entry_sure.delete(0, tk.END)
        self.listbox_ogr.selection_clear(0, tk.END)
        self.derslik_gerekmez_var.set(False)
        self.lab_gerekir_var.set(False)
        self.aktif_mi_var.set(True)
        self.combo_sabit_gun.set("(Esnek / Serbest)")
        self.combo_sabit_saat.set("(Esnek / Serbest)")

    def arayuz_ogretim_elemanlari(self):
        frame_sol = tk.Frame(self.tab_ogr)
        frame_sol.pack(side="left", fill="y", padx=10, pady=10)
        frame_sag = tk.Frame(self.tab_ogr)
        frame_sag.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        frame_ust = tk.Frame(frame_sol)
        frame_ust.pack(fill="x", pady=5)
        ttk.Label(frame_ust, text="Öğretim Elemanı Adı:").pack(side="left")
        self.entry_ogr_ad = ttk.Entry(frame_ust, width=25, font=self.font_standart)
        self.entry_ogr_ad.pack(side="right", fill="x", expand=True, padx=5)

        frame_matris = ttk.LabelFrame(frame_sol, text="Müsaitlik Matrisi", padding=10)
        frame_matris.pack(fill="x", pady=10)

        for j, gun in enumerate(self.gunler):
            tk.Button(frame_matris, text=gun, font=("Segoe UI", 8, "bold"), relief="groove", bg="#f0f0f0",
                      command=lambda g=gun: self.toggle_gun(g, self.ogr_musaitlik_vars)).grid(row=0, column=j+1, padx=2, pady=2, sticky="we")
            
        for i, saat in enumerate(self.saatler):
            row_idx = i + 1
            if i >= self.ogle_arasi_siniri: row_idx += 1
            if i == self.ogle_arasi_siniri:
                tk.Label(frame_matris, text="--- ÖĞLE ARASI ---", fg="red", font=("Segoe UI", 8, "bold")).grid(row=row_idx-1, column=0, columnspan=6)
            ttk.Label(frame_matris, text=saat, font=("Segoe UI", 8)).grid(row=row_idx, column=0, sticky="e", padx=5)
            for j, gun in enumerate(self.gunler):
                var = tk.BooleanVar(value=True)
                self.ogr_musaitlik_vars[(gun, saat)] = var
                ttk.Checkbutton(frame_matris, variable=var).grid(row=row_idx, column=j+1)

        frame_hizli = tk.Frame(frame_matris)
        frame_hizli.grid(row=len(self.saatler)+2, column=0, columnspan=6, pady=10)
        tk.Button(frame_hizli, text="Sabah", font=("Segoe UI", 8), command=lambda: self.toggle_blok(True, self.ogr_musaitlik_vars)).pack(side="left", padx=2)
        tk.Button(frame_hizli, text="Öğleden Sonra", font=("Segoe UI", 8), command=lambda: self.toggle_blok(False, self.ogr_musaitlik_vars)).pack(side="left", padx=2)
        tk.Button(frame_hizli, text="Tümünü Seç/Kaldır", font=("Segoe UI", 8), command=lambda: self.tumunu_sec(self.ogr_musaitlik_vars)).pack(side="left", padx=2)

        frame_liste_buton = tk.Frame(frame_sol)
        frame_liste_buton.pack(fill="x", pady=5)
        tk.Button(frame_liste_buton, text="Sisteme Ekle / Güncelle", command=self.ogr_ekle, bg="#4CAF50", fg="white", font=self.font_baslik, pady=5).pack(fill="x", pady=2)
        tk.Button(frame_liste_buton, text="Seçimi Temizle (Yeni)", command=self.form_sifirla_ogr, bg="#FFC107", font=self.font_standart).pack(fill="x", pady=2)

        ttk.Label(frame_sag, text="Kayıtlı Öğretim Elemanları", font=self.font_baslik).pack(anchor="w")
        self.tree_ogr = ttk.Treeview(frame_sag, columns=("Ad"), show="headings")
        self.tree_ogr.heading("Ad", text="Öğretim Elemanı")
        self.tree_ogr.pack(fill="both", expand=True, pady=5)
        self.tree_ogr.bind('<<TreeviewSelect>>', self.ogr_secildi)
        
        tk.Button(frame_sag, text="Seçili Olanı Sil", command=self.ogr_sil, bg="#F44336", fg="white", font=self.font_baslik).pack(anchor="e", pady=5)

    def arayuz_siniflar(self):
        frame_sol = tk.Frame(self.tab_siniflar)
        frame_sol.pack(side="left", fill="y", padx=10, pady=10)
        frame_sag = tk.Frame(self.tab_siniflar)
        frame_sag.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        frame_ust = tk.Frame(frame_sol)
        frame_ust.pack(fill="x", pady=5)
        ttk.Label(frame_ust, text="Mekan Adı:").pack(side="left")
        self.entry_sinif = ttk.Entry(frame_ust, width=20, font=self.font_standart)
        self.entry_sinif.pack(side="right", fill="x", expand=True, padx=5)

        frame_tip = tk.Frame(frame_sol)
        frame_tip.pack(fill="x", pady=5)
        ttk.Label(frame_tip, text="Mekan Türü:").pack(side="left")
        self.combo_sinif_tipi = ttk.Combobox(frame_tip, values=["Derslik", "Laboratuvar"], width=17, state="readonly", font=self.font_standart)
        self.combo_sinif_tipi.set("Derslik")
        self.combo_sinif_tipi.pack(side="right", padx=5)

        frame_matris = ttk.LabelFrame(frame_sol, text="Müsaitlik Matrisi", padding=10)
        frame_matris.pack(fill="x", pady=10)

        for j, gun in enumerate(self.gunler):
            tk.Button(frame_matris, text=gun, font=("Segoe UI", 8, "bold"), relief="groove", bg="#f0f0f0",
                      command=lambda g=gun: self.toggle_gun(g, self.sinif_musaitlik_vars)).grid(row=0, column=j+1, padx=2, pady=2, sticky="we")
            
        for i, saat in enumerate(self.saatler):
            row_idx = i + 1
            if i >= self.ogle_arasi_siniri: row_idx += 1
            if i == self.ogle_arasi_siniri:
                tk.Label(frame_matris, text="--- ÖĞLE ARASI ---", fg="red", font=("Segoe UI", 8, "bold")).grid(row=row_idx-1, column=0, columnspan=6)
            ttk.Label(frame_matris, text=saat, font=("Segoe UI", 8)).grid(row=row_idx, column=0, sticky="e", padx=5)
            for j, gun in enumerate(self.gunler):
                var = tk.BooleanVar(value=True)
                self.sinif_musaitlik_vars[(gun, saat)] = var
                ttk.Checkbutton(frame_matris, variable=var).grid(row=row_idx, column=j+1)

        frame_hizli = tk.Frame(frame_matris)
        frame_hizli.grid(row=len(self.saatler)+2, column=0, columnspan=6, pady=10)
        tk.Button(frame_hizli, text="Sabah", font=("Segoe UI", 8), command=lambda: self.toggle_blok(True, self.sinif_musaitlik_vars)).pack(side="left", padx=2)
        tk.Button(frame_hizli, text="Öğleden Sonra", font=("Segoe UI", 8), command=lambda: self.toggle_blok(False, self.sinif_musaitlik_vars)).pack(side="left", padx=2)
        tk.Button(frame_hizli, text="Tümünü Seç/Kaldır", font=("Segoe UI", 8), command=lambda: self.tumunu_sec(self.sinif_musaitlik_vars)).pack(side="left", padx=2)

        frame_liste_buton = tk.Frame(frame_sol)
        frame_liste_buton.pack(fill="x", pady=5)
        tk.Button(frame_liste_buton, text="Sisteme Ekle / Güncelle", command=self.sinif_ekle, bg="#4CAF50", fg="white", font=self.font_baslik, pady=5).pack(fill="x", pady=2)
        tk.Button(frame_liste_buton, text="Seçimi Temizle (Yeni)", command=self.form_sifirla_sinif, bg="#FFC107", font=self.font_standart).pack(fill="x", pady=2)

        ttk.Label(frame_sag, text="Kayıtlı Sınıf ve Laboratuvarlar", font=self.font_baslik).pack(anchor="w")
        self.tree_sinif = ttk.Treeview(frame_sag, columns=("Ad", "Tip"), show="headings")
        self.tree_sinif.heading("Ad", text="Mekan Adı")
        self.tree_sinif.heading("Tip", text="Tür")
        self.tree_sinif.column("Ad", width=180)
        self.tree_sinif.column("Tip", width=100, anchor="center")
        self.tree_sinif.pack(fill="both", expand=True, pady=5)
        self.tree_sinif.bind('<<TreeviewSelect>>', self.sinif_secildi)
        
        tk.Button(frame_sag, text="Seçili Olanı Sil", command=self.sinif_sil, bg="#F44336", fg="white", font=self.font_baslik).pack(anchor="e", pady=5)

    def arayuz_dersler(self):
        frame_sol = tk.Frame(self.tab_dersler)
        frame_sol.pack(side="left", fill="y", padx=10, pady=10)
        frame_sag = tk.Frame(self.tab_dersler)
        frame_sag.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        frame_giris = ttk.LabelFrame(frame_sol, text="Ders ve Müfredat Bilgileri", padding=10)
        frame_giris.pack(fill="x")

        ttk.Label(frame_giris, text="Dersin Adı:").grid(row=0, column=0, sticky="nw", pady=5)
        self.entry_ders = ttk.Entry(frame_giris, width=30, font=self.font_standart)
        self.entry_ders.grid(row=0, column=1, padx=5, pady=5)

        ttk.Label(frame_giris, text="Öğr. Elemanı:\n(CTRL ile çoklu)").grid(row=1, column=0, sticky="nw", pady=5)
        frame_listbox_ogr = tk.Frame(frame_giris)
        frame_listbox_ogr.grid(row=1, column=1, padx=5, pady=5, sticky="we")
        
        self.listbox_ogr = tk.Listbox(frame_listbox_ogr, selectmode=tk.MULTIPLE, height=4, font=self.font_standart, exportselection=False)
        scrollbar_ogr = ttk.Scrollbar(frame_listbox_ogr, orient="vertical", command=self.listbox_ogr.yview)
        self.listbox_ogr.configure(yscrollcommand=scrollbar_ogr.set)
        self.listbox_ogr.pack(side="left", fill="both", expand=True)
        scrollbar_ogr.pack(side="right", fill="y")

        ttk.Label(frame_giris, text="Süre (Blok):").grid(row=2, column=0, sticky="nw", pady=5)
        self.entry_sure = ttk.Entry(frame_giris, width=30, font=self.font_standart)
        self.entry_sure.grid(row=2, column=1, padx=5, pady=5)

        ttk.Label(frame_giris, text="Sınıf Seviyesi:").grid(row=3, column=0, sticky="nw", pady=5)
        self.combo_sinif_yili = ttk.Combobox(frame_giris, values=["1. Sınıf", "2. Sınıf", "3. Sınıf", "4. Sınıf", "Seçmeli/Karma"], width=28, state="readonly", font=self.font_standart)
        self.combo_sinif_yili.set("1. Sınıf")
        self.combo_sinif_yili.grid(row=3, column=1, padx=5, pady=5)

        ttk.Label(frame_giris, text="Dönem / Yıl:").grid(row=4, column=0, sticky="nw", pady=5)
        frame_donem = tk.Frame(frame_giris)
        frame_donem.grid(row=4, column=1, padx=5, pady=5, sticky="we")
        self.combo_donem_kayit = ttk.Combobox(frame_donem, values=["Güz", "Bahar", "Yaz"], width=10, state="readonly", font=self.font_standart)
        self.combo_donem_kayit.set("Güz")
        self.combo_donem_kayit.pack(side="left", padx=(0,5))
        self.combo_mufredat = ttk.Combobox(frame_donem, values=["2023", "2024", "2026", "Ortak"], width=10, state="readonly", font=self.font_standart)
        self.combo_mufredat.set("2026")
        self.combo_mufredat.pack(side="left")

        # ÖZEL AYARLAR
        frame_kurallar = ttk.LabelFrame(frame_sol, text="Özel Ayarlar", padding=10)
        frame_kurallar.pack(fill="x", pady=5)
        ttk.Checkbutton(frame_kurallar, text="Ders Bu Dönem AÇILACAK (Aktif)", variable=self.aktif_mi_var).pack(anchor="w", pady=2)
        ttk.Checkbutton(frame_kurallar, text="Derslik Gerektirmez (Ofis/Dış Uygulama)", variable=self.derslik_gerekmez_var).pack(anchor="w", pady=2)
        ttk.Checkbutton(frame_kurallar, text="Laboratuvar Ortamı Gerekir", variable=self.lab_gerekir_var).pack(anchor="w", pady=2)

        # SABİT GÜN / SAAT KİLİDİ
        frame_sabit = ttk.LabelFrame(frame_sol, text="Sabit Gün / Saat Kilidi (Opsiyonel)", padding=10)
        frame_sabit.pack(fill="x", pady=5)

        ttk.Label(frame_sabit, text="Sabit Gün:").grid(row=0, column=0, sticky="w", pady=2)
        self.combo_sabit_gun = ttk.Combobox(frame_sabit, values=["(Esnek / Serbest)"] + self.gunler, width=20, state="readonly", font=self.font_standart)
        self.combo_sabit_gun.set("(Esnek / Serbest)")
        self.combo_sabit_gun.grid(row=0, column=1, sticky="w", pady=2, padx=5)

        ttk.Label(frame_sabit, text="Sabit Başl. Saati:").grid(row=1, column=0, sticky="w", pady=2)
        self.combo_sabit_saat = ttk.Combobox(frame_sabit, values=["(Esnek / Serbest)"] + self.saatler, width=20, state="readonly", font=self.font_standart)
        self.combo_sabit_saat.set("(Esnek / Serbest)")
        self.combo_sabit_saat.grid(row=1, column=1, sticky="w", pady=2, padx=5)

        tk.Button(frame_sol, text="Havuza Ekle / Güncelle", command=self.ders_ekle, bg="#4CAF50", fg="white", font=self.font_baslik, pady=5).pack(fill="x", pady=5)
        tk.Button(frame_sol, text="Seçimi Temizle (Yeni)", command=self.form_sifirla_ders, bg="#FFC107", font=self.font_standart).pack(fill="x")

        ttk.Label(frame_sag, text="Veritabanındaki Dersler (Sekmeler Arası Geçiş Yapabilirsiniz)", font=self.font_baslik).pack(anchor="w", pady=(0, 5))
        
        self.notebook_mufredat = ttk.Notebook(frame_sag)
        self.notebook_mufredat.pack(fill="both", expand=True)
        self.notebook_mufredat.bind("<<NotebookTabChanged>>", lambda e: self.form_sifirla_ders())

        frame_2026 = ttk.Frame(self.notebook_mufredat)
        frame_2024 = ttk.Frame(self.notebook_mufredat)
        frame_diger = ttk.Frame(self.notebook_mufredat)

        self.notebook_mufredat.add(frame_2026, text="   2026 Müfredatı   ")
        self.notebook_mufredat.add(frame_2024, text="   2024 Müfredatı   ")
        self.notebook_mufredat.add(frame_diger, text="   Diğer (2023/Ortak)   ")

        self.tree_2026 = self.ders_tablosu_olustur(frame_2026)
        self.tree_2024 = self.ders_tablosu_olustur(frame_2024)
        self.tree_diger = self.ders_tablosu_olustur(frame_diger)

        frame_sag_alt = tk.Frame(frame_sag)
        frame_sag_alt.pack(fill="x", pady=5)
        tk.Button(frame_sag_alt, text="Seçili Dersleri AKTİF YAP", command=lambda: self.ders_toplu_durum(True), bg="#8BC34A", font=self.font_baslik).pack(side="left", padx=2)
        tk.Button(frame_sag_alt, text="Seçili Dersleri PASİF YAP", command=lambda: self.ders_toplu_durum(False), bg="#FFB74D", font=self.font_baslik).pack(side="left", padx=2)
        tk.Button(frame_sag_alt, text="Seçili Dersi SİL", command=self.ders_sil, bg="#F44336", fg="white", font=self.font_baslik).pack(side="right", padx=2)

    def ders_tablosu_olustur(self, parent):
        tree = ttk.Treeview(parent, columns=("Durum", "Donem", "Ders", "Sinif", "Ogr", "Sure", "Mekan", "Sabit"), show="headings", selectmode="extended")
        tree.heading("Durum", text="Durum")
        tree.heading("Donem", text="Dönem")
        tree.heading("Ders", text="Ders Adı")
        tree.heading("Sinif", text="Seviye")
        tree.heading("Ogr", text="Öğr. Elemanı")
        tree.heading("Sure", text="Süre")
        tree.heading("Mekan", text="Mekan/Tip")
        tree.heading("Sabit", text="Sabit Kilit")
        
        tree.column("Durum", width=55, anchor="center")
        tree.column("Donem", width=45, anchor="center")
        tree.column("Ders", width=140)
        tree.column("Sinif", width=65, anchor="center")
        tree.column("Ogr", width=120)
        tree.column("Sure", width=35, anchor="center")
        tree.column("Mekan", width=75, anchor="center")
        tree.column("Sabit", width=85, anchor="center")
        
        tree.pack(fill="both", expand=True, padx=2, pady=2)
        tree.bind('<<TreeviewSelect>>', self.ders_secildi)
        return tree

    def aktif_tabloyu_getir(self):
        tab_id = self.notebook_mufredat.index("current")
        if tab_id == 0: return self.tree_2026
        elif tab_id == 1: return self.tree_2024
        else: return self.tree_diger

    def arayuz_sonuc(self):
        frame_ust = tk.Frame(self.tab_sonuc)
        frame_ust.pack(fill="x", padx=10, pady=10)

        frame_donem = ttk.LabelFrame(frame_ust, text="Program Ayarları", padding=10)
        frame_donem.pack(side="left", fill="x", expand=True)

        ttk.Label(frame_donem, text="Hangi Dönem İçin Program Üretilecek?").pack(side="left", padx=5)
        self.combo_donem_secim = ttk.Combobox(frame_donem, values=["Güz", "Bahar", "Yaz"], state="readonly", width=15, font=self.font_standart)
        self.combo_donem_secim.set("Güz")
        self.combo_donem_secim.pack(side="left", padx=10)

        tk.Button(frame_ust, text="Programı Üret (Yapay Zeka)", command=self.program_uret, bg="#2196F3", fg="white", font=("Segoe UI", 12, "bold"), padx=15).pack(side="left", padx=10)
        tk.Button(frame_ust, text="Excel'e Aktar", command=self.excel_aktar, bg="#FF9800", fg="white", font=("Segoe UI", 12, "bold"), padx=15).pack(side="left", padx=10)

        self.tree_sonuc = ttk.Treeview(self.tab_sonuc, columns=("Gun", "Saat", "Derslik", "Ders", "Sinif", "Ogr"), show="headings")
        self.tree_sonuc.heading("Gun", text="Gün")
        self.tree_sonuc.heading("Saat", text="Saat")
        self.tree_sonuc.heading("Derslik", text="Derslik/Mekan")
        self.tree_sonuc.heading("Ders", text="Ders Adı")
        self.tree_sonuc.heading("Sinif", text="Seviye")
        self.tree_sonuc.heading("Ogr", text="Öğr. Elemanı")
        
        self.tree_sonuc.column("Gun", width=90, anchor="center")
        self.tree_sonuc.column("Saat", width=90, anchor="center")
        self.tree_sonuc.column("Derslik", width=120, anchor="center")
        self.tree_sonuc.column("Ders", width=200)
        self.tree_sonuc.column("Sinif", width=80, anchor="center")
        self.tree_sonuc.column("Ogr", width=180)

        self.tree_sonuc.pack(padx=10, pady=10, fill="both", expand=True)

    def ogr_secildi(self, event):
        secili = self.tree_ogr.selection()
        if not secili: return
        idx = self.tree_ogr.index(secili[0])
        o = self.ogretim_elemanlari[idx]
        self.entry_ogr_ad.delete(0, tk.END)
        self.entry_ogr_ad.insert(0, o["ad"])
        for (g, s), v in self.ogr_musaitlik_vars.items():
            v.set(o["musaitlik"].get((g, s), True))

    def sinif_secildi(self, event):
        secili = self.tree_sinif.selection()
        if not secili: return
        idx = self.tree_sinif.index(secili[0])
        s = self.siniflar[idx]
        self.entry_sinif.delete(0, tk.END)
        self.entry_sinif.insert(0, s["ad"])
        self.combo_sinif_tipi.set(s.get("tip", "Derslik"))
        for (g, st), v in self.sinif_musaitlik_vars.items():
            v.set(s["musaitlik"].get((g, st), True))

    def ders_secildi(self, event):
        aktif_agac = self.aktif_tabloyu_getir()
        secili = aktif_agac.selection()
        if not secili: return
        
        idx = int(aktif_agac.item(secili[0], "tags")[0])
        d = self.dersler[idx]
        
        self.entry_ders.delete(0, tk.END)
        self.entry_ders.insert(0, d["id"])
        
        self.listbox_ogr.selection_clear(0, tk.END)
        for i, h in enumerate(self.ogretim_elemanlari):
            if h["ad"] in d["hocalar"]: self.listbox_ogr.selection_set(i)
                
        self.entry_sure.delete(0, tk.END)
        self.entry_sure.insert(0, str(d["sure"]))
        self.combo_sinif_yili.set(d.get("sinif_yil", "1. Sınıf"))
        self.combo_mufredat.set(d.get("mufredat", "2026"))
        self.combo_donem_kayit.set(d.get("donem", "Güz"))
        self.derslik_gerekmez_var.set(d.get("derslik_gerekmez", False))
        self.lab_gerekir_var.set(d.get("lab_gerekir", False))
        self.aktif_mi_var.set(d.get("aktif_mi", True))
        self.combo_sabit_gun.set(d.get("sabit_gun", "(Esnek / Serbest)"))
        self.combo_sabit_saat.set(d.get("sabit_saat", "(Esnek / Serbest)"))

    def veri_kaydet(self):
        kayit_verisi = {
            "bolum_adi": self.bolum_adi,
            "ogretim_elemanlari": [
                {"ad": o["ad"], "musaitlik": {f"{g}_{s}": v for (g, s), v in o["musaitlik"].items()}} for o in self.ogretim_elemanlari
            ],
            "siniflar": [
                {"ad": s["ad"], "tip": s.get("tip", "Derslik"), "musaitlik": {f"{g}_{s}": v for (g, s), v in s["musaitlik"].items()}} for s in self.siniflar
            ],
            "dersler": self.dersler
        }
        with open(self.veri_dosyasi, "w", encoding="utf-8") as f:
            json.dump(kayit_verisi, f, ensure_ascii=False, indent=4)

    def veri_yukle(self):
        if not os.path.exists(self.veri_dosyasi): return
        try:
            with open(self.veri_dosyasi, "r", encoding="utf-8") as f:
                kayit_verisi = json.load(f)

            self.bolum_adi = kayit_verisi.get("bolum_adi", "BÖLÜMÜ")
            self.root.title(f"KMÜ {self.bolum_adi} Ders Programı Planlayıcısı")

            for o in kayit_verisi.get("ogretim_elemanlari", []):
                musaitlik = {tuple(k.split("_")): v for k, v in o["musaitlik"].items()}
                self.ogretim_elemanlari.append({"ad": o["ad"], "musaitlik": musaitlik})
                self.tree_ogr.insert("", tk.END, values=(o["ad"],))

            for s in kayit_verisi.get("siniflar", []):
                musaitlik = {tuple(k.split("_")): v for k, v in s["musaitlik"].items()}
                s_tip = s.get("tip", "Derslik")
                self.siniflar.append({"ad": s["ad"], "tip": s_tip, "musaitlik": musaitlik})
                self.tree_sinif.insert("", tk.END, values=(s["ad"], s_tip))

            for d in kayit_verisi.get("dersler", []):
                hoca_listesi = d.get("hocalar", d.get("hoca", ["Atanmadı"]))
                if isinstance(hoca_listesi, str): hoca_listesi = [hoca_listesi]
                
                d_dict = {
                    "id": d["id"], "hocalar": hoca_listesi, "sure": d["sure"],
                    "sinif_yil": d.get("sinif_yil", "Belirtilmemiş"), "mufredat": d.get("mufredat", "2026"), 
                    "donem": d.get("donem", "Güz"), "derslik_gerekmez": d.get("derslik_gerekmez", False),
                    "lab_gerekir": d.get("lab_gerekir", False), "aktif_mi": d.get("aktif_mi", True),
                    "sabit_gun": d.get("sabit_gun", "(Esnek / Serbest)"), "sabit_saat": d.get("sabit_saat", "(Esnek / Serbest)")
                }
                self.dersler.append(d_dict)

            self.listbox_guncelle()
            self.ders_tablolarini_yenile()

        except Exception as e:
            messagebox.showerror("Hata", f"Veriler yüklenemedi:\n{str(e)}")

    def listbox_guncelle(self):
        self.listbox_ogr.delete(0, tk.END)
        for o in self.ogretim_elemanlari:
            self.listbox_ogr.insert(tk.END, o["ad"])

    def ders_tablolarini_yenile(self):
        for tree in [self.tree_2026, self.tree_2024, self.tree_diger]:
            for item in tree.get_children(): tree.delete(item)

        for i, d in enumerate(self.dersler):
            durum_metni = "✅ Aktif" if d["aktif_mi"] else "❌ Pasif"
            hoca_metni = ", ".join(d["hocalar"])
            if d["derslik_gerekmez"]:
                mekan_metni = "Ofis/Dış"
            elif d.get("lab_gerekir", False):
                mekan_metni = "Laboratuvar"
            else:
                mekan_metni = "Derslik"
            
            sabit_gun_str = d.get('sabit_gun', '(Esnek / Serbest)')
            sabit_saat_str = d.get('sabit_saat', '(Esnek / Serbest)')
            
            # --- YAPILAN BÜYÜK DÜZELTME BURADA (DEĞİŞKEN ADI HATASI GİDERİLDİ) ---
            if sabit_gun_str != "(Esnek / Serbest)" and sabit_saat_str != "(Esnek / Serbest)":
                sabit_metni = f"{sabit_gun_str[:3]} {sabit_saat_str[:5]}"
            elif sabit_gun_str != "(Esnek / Serbest)":
                sabit_metni = f"{sabit_gun_str[:3]} (Serb. Saat)"
            elif sabit_saat_str != "(Esnek / Serbest)":
                sabit_metni = f"Her Gün {sabit_saat_str[:5]}"
            else:
                sabit_metni = "Serbest"
            
            degerler = (durum_metni, d["donem"], d["id"], d["sinif_yil"], hoca_metni, d["sure"], mekan_metni, sabit_metni)
            
            if d["mufredat"] == "2026":
                self.tree_2026.insert("", tk.END, values=degerler, tags=(str(i),))
            elif d["mufredat"] == "2024":
                self.tree_2024.insert("", tk.END, values=degerler, tags=(str(i),))
            else:
                self.tree_diger.insert("", tk.END, values=degerler, tags=(str(i),))

    def ogr_ekle(self):
        ad = self.entry_ogr_ad.get().strip()
        if not ad: return messagebox.showwarning("Uyarı", "Ad boş olamaz.")
        musaitlik_durumu = {(g, s): self.ogr_musaitlik_vars[(g, s)].get() for g in self.gunler for s in self.saatler}
        secili = self.tree_ogr.selection()
        if secili:
            idx = self.tree_ogr.index(secili[0])
            self.ogretim_elemanlari[idx] = {"ad": ad, "musaitlik": musaitlik_durumu}
            self.tree_ogr.item(secili[0], values=(ad,))
            self.listbox_guncelle()
        else:
            if any(o["ad"] == ad for o in self.ogretim_elemanlari): return messagebox.showwarning("Uyarı", "Bu isim zaten var.")
            self.ogretim_elemanlari.append({"ad": ad, "musaitlik": musaitlik_durumu})
            self.tree_ogr.insert("", tk.END, values=(ad,))
            self.listbox_guncelle()
        self.veri_kaydet()

    def ogr_sil(self):
        secili = self.tree_ogr.selection()
        if not secili: return
        idx = self.tree_ogr.index(secili[0])
        del self.ogretim_elemanlari[idx]
        self.tree_ogr.delete(secili[0])
        self.listbox_guncelle()
        self.veri_kaydet()
        self.form_sifirla_ogr()

    def sinif_ekle(self):
        ad = self.entry_sinif.get().strip()
        s_tip = self.combo_sinif_tipi.get()
        if not ad: return messagebox.showwarning("Uyarı", "Mekan adı boş olamaz.")
        musaitlik_durumu = {(g, s): self.sinif_musaitlik_vars[(g, s)].get() for g in self.gunler for s in self.saatler}
        secili = self.tree_sinif.selection()
        if secili:
            idx = self.tree_sinif.index(secili[0])
            self.siniflar[idx] = {"ad": ad, "tip": s_tip, "musaitlik": musaitlik_durumu}
            self.tree_sinif.item(secili[0], values=(ad, s_tip))
        else:
            if any(s["ad"] == ad for s in self.siniflar): return messagebox.showwarning("Uyarı", "Bu mekan zaten var.")
            self.siniflar.append({"ad": ad, "tip": s_tip, "musaitlik": musaitlik_durumu})
            self.tree_sinif.insert("", tk.END, values=(ad, s_tip))
        self.veri_kaydet()

    def sinif_sil(self):
        secili = self.tree_sinif.selection()
        if not secili: return
        idx = self.tree_sinif.index(secili[0])
        del self.siniflar[idx]
        self.tree_sinif.delete(secili[0])
        self.veri_kaydet()
        self.form_sifirla_sinif()

    def ders_toplu_durum(self, durum):
        aktif_agac = self.aktif_tabloyu_getir()
        secililer = aktif_agac.selection()
        if not secililer: return messagebox.showwarning("Uyarı", "Lütfen tablodan en az bir ders seçin.")
        for item in secililer:
            idx = int(aktif_agac.item(item, "tags")[0])
            self.dersler[idx]["aktif_mi"] = durum
        self.veri_kaydet()
        self.ders_tablolarini_yenile()

    def ders_ekle(self):
        ad = self.entry_ders.get().strip()
        secili_hocalar_idx = self.listbox_ogr.curselection()
        if not secili_hocalar_idx: return messagebox.showwarning("Uyarı", "En az bir Öğr. Elemanı seçin.")
        secilen_hocalar = [self.listbox_ogr.get(i) for i in secili_hocalar_idx]
        
        yeni_ders = {
            "id": ad, "hocalar": secilen_hocalar, 
            "sinif_yil": self.combo_sinif_yili.get(), "mufredat": self.combo_mufredat.get(), 
            "donem": self.combo_donem_kayit.get(), "derslik_gerekmez": self.derslik_gerekmez_var.get(), 
            "lab_gerekir": self.lab_gerekir_var.get(), "aktif_mi": self.aktif_mi_var.get(),
            "sabit_gun": self.combo_sabit_gun.get(), "sabit_saat": self.combo_sabit_saat.get()
        }
        
        try:
            yeni_ders["sure"] = int(self.entry_sure.get().strip())
            if not ad: raise ValueError
            
            aktif_agac = self.aktif_tabloyu_getir()
            secili = aktif_agac.selection()
            
            if len(secili) == 1:
                idx = int(aktif_agac.item(secili[0], "tags")[0])
                self.dersler[idx] = yeni_ders
            else:
                self.dersler.append(yeni_ders)
                
            self.veri_kaydet()
            self.ders_tablolarini_yenile()
        except:
            messagebox.showerror("Hata", "Eksik/hatalı giriş. (Süreye rakam yazınız).")

    def ders_sil(self):
        aktif_agac = self.aktif_tabloyu_getir()
        secililer = aktif_agac.selection()
        if not secililer: return
        indeksler = sorted([int(aktif_agac.item(item, "tags")[0]) for item in secililer], reverse=True)
        for idx in indeksler:
            del self.dersler[idx]
        self.veri_kaydet()
        self.form_sifirla_ders()
        self.ders_tablolarini_yenile()

    def program_uret(self):
        secilen_donem = self.combo_donem_secim.get()
        aktif_dersler = [d for d in self.dersler if d["donem"] == secilen_donem and d.get("aktif_mi", True)]
        
        if not aktif_dersler: return messagebox.showwarning("Uyarı", f"'{secilen_donem}' döneminde 'Aktif' olarak işaretlenmiş ders bulunamadı.")
        if not self.siniflar or not self.ogretim_elemanlari: return messagebox.showwarning("Uyarı", "Sınıf/Laboratuvar veya öğretim elemanı eklemediniz.")

        onay_mesaji = f"Seçilen Dönem: {secilen_donem}\n\nAşağıdaki AKTİF dersler için program yapılacaktır:\n\n"
        for d in aktif_dersler: onay_mesaji += f"• {d['id']} ({', '.join(d['hocalar'])}) - {d['sure']} Saat\n"
        onay_mesaji += "\nOnaylıyor musunuz?"
        if not messagebox.askyesno("Onay", onay_mesaji): return

        for item in self.tree_sonuc.get_children(): self.tree_sonuc.delete(item)
        self.olusturulan_program = []
        model = cp_model.CpModel()
        gun_sayisi, gunluk_saat = len(self.gunler), len(self.saatler)
        
        ders_baslangic = {}
        ogr_araliklari = {o["ad"]: [] for o in self.ogretim_elemanlari}
        sinif_araliklari = {s["ad"]: [] for s in self.siniflar}
        sinif_secimleri = {}
        
        ogrenci_araliklari = {"1. Sınıf": [], "2. Sınıf": [], "3. Sınıf": [], "4. Sınıf": [], "Seçmeli/Karma": []}
        
        for o in self.ogretim_elemanlari:
            for gun_idx, gun in enumerate(self.gunler):
                for saat_idx, saat in enumerate(self.saatler):
                    if not o["musaitlik"][(gun, saat)]:
                        zaman = gun_idx * gunluk_saat + saat_idx
                        ogr_araliklari[o["ad"]].append(model.NewIntervalVar(zaman, 1, zaman + 1, f'ko_{o["ad"]}_{zaman}'))

        for s in self.siniflar:
            for gun_idx, gun in enumerate(self.gunler):
                for saat_idx, saat in enumerate(self.saatler):
                    if not s["musaitlik"][(gun, saat)]:
                        zaman = gun_idx * gunluk_saat + saat_idx
                        sinif_araliklari[s["ad"]].append(model.NewIntervalVar(zaman, 1, zaman + 1, f'ks_{s["ad"]}_{zaman}'))

        for i, ders in enumerate(aktif_dersler):
            sure = ders["sure"]
            gecerli_saatler = []
            
            for hoca in ders["hocalar"]:
                if hoca not in ogr_araliklari: return messagebox.showerror("Hata", f"Hoca bulunamadı: {hoca}")

            sabit_gun = ders.get("sabit_gun", "(Esnek / Serbest)")
            sabit_saat = ders.get("sabit_saat", "(Esnek / Serbest)")

            if sabit_gun == "(Esnek / Serbest)" and sabit_saat == "(Esnek / Serbest)":
                for gun_idx in range(gun_sayisi):
                    for saat_idx in range(gunluk_saat):
                        if sure == 8:
                            if saat_idx == 0: gecerli_saatler.append(gun_idx * gunluk_saat + 0)
                        else:
                            if saat_idx + sure > gunluk_saat: continue
                            if saat_idx < self.ogle_arasi_siniri and (saat_idx + sure) > self.ogle_arasi_siniri: continue
                            gecerli_saatler.append(gun_idx * gunluk_saat + saat_idx)
            else:
                if sabit_gun != "(Esnek / Serbest)" and sabit_saat != "(Esnek / Serbest)":
                    g_idx = self.gunler.index(sabit_gun)
                    s_idx = self.saatler.index(sabit_saat)
                    if sure == 8 and s_idx != 0:
                        return messagebox.showerror("Hata", f"'{ders['id']}' 8 saatlik ders olduğu için saat 08:30 olmalıdır.")
                    if sure != 8 and (s_idx + sure > gunluk_saat or (s_idx < self.ogle_arasi_siniri and (s_idx + sure) > self.ogle_arasi_siniri)):
                        return messagebox.showerror("Hata", f"'{ders['id']}' sabit saati öğle arasına veya mesai dışına taşıyor!")
                    gecerli_saatler.append(g_idx * gunluk_saat + s_idx)
                    
                elif sabit_gun != "(Esnek / Serbest)" and sabit_saat == "(Esnek / Serbest)":
                    g_idx = self.gunler.index(sabit_gun)
                    for saat_idx in range(gunluk_saat):
                        if sure == 8:
                            if saat_idx == 0: gecerli_saatler.append(g_idx * gunluk_saat + 0)
                        else:
                            if saat_idx + sure > gunluk_saat: continue
                            if saat_idx < self.ogle_arasi_siniri and (saat_idx + sure) > self.ogle_arasi_siniri: continue
                            gecerli_saatler.append(g_idx * gunluk_saat + saat_idx)
                            
                elif sabit_gun == "(Esnek / Serbest)" and sabit_saat != "(Esnek / Serbest)":
                    s_idx = self.saatler.index(sabit_saat)
                    if sure == 8 and s_idx != 0:
                        return messagebox.showerror("Hata", f"'{ders['id']}' 8 saatlik ders olduğu için saat 08:30 olmalıdır.")
                    if sure != 8 and (s_idx + sure > gunluk_saat or (s_idx < self.ogle_arasi_siniri and (s_idx + sure) > self.ogle_arasi_siniri)):
                        return messagebox.showerror("Hata", f"'{ders['id']}' sabit saati öğle arasına veya mesai dışına taşıyor!")
                    for g_idx in range(gun_sayisi):
                        gecerli_saatler.append(g_idx * gunluk_saat + s_idx)

            if not gecerli_saatler: return messagebox.showerror("Hata", f"Ders sığmıyor veya kurallara uymuyor: {ders['id']}")

            baslangic = model.NewIntVarFromDomain(cp_model.Domain.FromValues(gecerli_saatler), f'bas_{i}')
            bitis = model.NewIntVar(0, gun_sayisi * gunluk_saat, f'bitis_{i}')
            aralik = model.NewIntervalVar(baslangic, sure, bitis, f'aralik_{i}')
            
            ders_baslangic[i] = baslangic
            for hoca in ders["hocalar"]: ogr_araliklari[hoca].append(aralik)

            sy = ders.get("sinif_yil", "1. Sınıf")
            if sy in ogrenci_araliklari:
                ogrenci_araliklari[sy].append(aralik)

            if not ders.get("derslik_gerekmez", False):
                ders_sinif_secimleri = []
                lab_lazim = ders.get("lab_gerekir", False)
                
                for s in self.siniflar:
                    mekan_tipi = s.get("tip", "Derslik")
                    if lab_lazim and mekan_tipi != "Laboratuvar": continue
                    if not lab_lazim and mekan_tipi == "Laboratuvar": continue

                    secildi_mi = model.NewBoolVar(f'sec_{i}_{s["ad"]}')
                    ders_sinif_secimleri.append(secildi_mi)
                    sinif_secimleri[(i, s["ad"])] = secildi_mi
                    sinif_araliklari[s["ad"]].append(model.NewOptionalIntervalVar(baslangic, sure, bitis, secildi_mi, f'ops_{i}_{s["ad"]}'))
                
                if not ders_sinif_secimleri:
                    lab_uyari = " Laboratuvar" if lab_lazim else " Derslik"
                    return messagebox.showerror("Mekan Hatası", f"'{ders['id']}' için uygun türde ({lab_uyari}) kayıtlı mekan bulunamadı!")
                
                model.AddExactlyOne(ders_sinif_secimleri)

        for araliklar in ogr_araliklari.values(): model.AddNoOverlap(araliklar)
        for araliklar in sinif_araliklari.values(): model.AddNoOverlap(araliklar)
        for araliklar in ogrenci_araliklari.values(): 
            if araliklar: model.AddNoOverlap(araliklar)

        self.root.update()
        solver = cp_model.CpSolver()
        solver.parameters.max_time_in_seconds = 20.0 
        
        status = solver.Solve(model)

        if status == cp_model.UNKNOWN:
            messagebox.showerror("Zamanaşımı / Darboğaz", "Yapay zeka 20 saniye boyunca denedi ancak bu dar kısıtlamalarla bir program çıkaramadı.\n\nLütfen sabit gün/saat kilitlerini, mekan tiplerini ve hoca müsaitliklerini kontrol edin.")
            return

        elif status == cp_model.OPTIMAL or status == cp_model.FEASIBLE:
            for i, ders in enumerate(aktif_dersler):
                bas = solver.Value(ders_baslangic[i])
                sure = ders["sure"]
                
                if ders.get("derslik_gerekmez", False):
                    atanan_sinif = "Belirtilmedi (Ofis/Dış)"
                else:
                    atanan_sinif = next((s["ad"] for s in self.siniflar if (i, s["ad"]) in sinif_secimleri and solver.Value(sinif_secimleri[(i, s["ad"])]) == 1), "Atanmadı")
                
                hoca_metni = ", ".join(ders["hocalar"])
                
                for saat_artisi in range(sure):
                    gecerli_zaman = bas + saat_artisi
                    gun = self.gunler[gecerli_zaman // gunluk_saat]
                    saat = self.saatler[gecerli_zaman % gunluk_saat]
                    
                    self.olusturulan_program.append({
                        "Sıra (Gizli)": gecerli_zaman, "Dönem": ders["donem"], "Müfredat": ders["mufredat"],
                        "Sınıf Seviyesi": ders["sinif_yil"], "Gün": gun, "Saat": saat, "Derslik": atanan_sinif,
                        "Ders": ders["id"], "Öğretim Elemanı": hoca_metni, "Süre": f"{ders['sure']} Saat"
                    })
            
            self.olusturulan_program.sort(key=lambda x: x["Sıra (Gizli)"])
            for v in self.olusturulan_program:
                self.tree_sonuc.insert("", tk.END, values=(v["Gün"], v["Saat"], v["Derslik"], v["Ders"], v["Sınıf Seviyesi"], v["Öğretim Elemanı"]))
            messagebox.showinfo("Başarılı", f"{secilen_donem} ders programı kısıtlara uygun olarak üretildi!")
        else:
            messagebox.showerror("Çözüm Bulunamadı", "Saatler ve ders yükü matematiksel olarak çakışıyor, program yerleştirilemedi. Sabit gün kilitlerinin çakışıp çakışmadığını kontrol edin.")

    def excel_aktar(self):
        if not self.olusturulan_program: 
            return messagebox.showwarning("Uyarı", "Önce programı üretmelisiniz.")
            
        dosya = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], title="Kaydet")
        if not dosya: return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, Border, Side

            wb = Workbook()
            ws = wb.active
            ws.title = "Ders Programı"
            
            sinif_seviyeleri = ["1. Sınıf", "2. Sınıf", "3. Sınıf", "4. Sınıf", "Seçmeli/Karma"]
            
            ws.merge_cells('A1:G1')
            baslik = ws['A1']
            baslik.value = f"KARAMANOĞLU MEHMETBEY ÜNİVERSİTESİ {self.bolum_adi} N.Ö"
            baslik.font = Font(bold=True, size=12)
            baslik.alignment = Alignment(horizontal="center", vertical="center")
            
            headers = ["GÜN", "SAAT"] + sinif_seviyeleri
            ws.append(headers)
            
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=2, column=col_num)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            matris = { (g, s): {lvl: "" for lvl in sinif_seviyeleri} for g in self.gunler for s in self.saatler }
            
            for p in self.olusturulan_program:
                g = p["Gün"]
                s = p["Saat"]
                lvl = p["Sınıf Seviyesi"]
                
                ders_ad = p["Ders"]
                if " " in ders_ad and ders_ad.split(" ")[0].isdigit():
                    ders_ad = " ".join(ders_ad.split(" ")[1:])
                    
                metin = f"{ders_ad}\n({p['Öğretim Elemanı']})\n{p['Derslik']}"
                
                if matris[(g, s)][lvl] == "":
                    matris[(g, s)][lvl] = metin
                else:
                    matris[(g, s)][lvl] += f"\n\n{metin}" 
                    
            border_style = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            
            row_idx = 3
            for g in self.gunler:
                start_row = row_idx
                for s in self.saatler:
                    row_data = [g, s] + [matris[(g, s)][lvl] for lvl in sinif_seviyeleri]
                    ws.append(row_data)
                    ws.row_dimensions[row_idx].height = 65
                    
                    for col_idx in range(1, 8):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        cell.border = border_style
                    row_idx += 1
                    
                ws.merge_cells(start_row=start_row, start_column=1, end_row=row_idx-1, end_column=1)
                ws.cell(row=start_row, column=1).alignment = Alignment(horizontal="center", vertical="center", textRotation=90)
                
            ws.column_dimensions['A'].width = 6
            ws.column_dimensions['B'].width = 15
            for c in ['C', 'D', 'E', 'F', 'G']:
                ws.column_dimensions[c].width = 25
                
            wb.save(dosya)
            messagebox.showinfo("Başarılı", f"Program, okulunuzun formatında Excel'e aktarıldı:\n{dosya}")

        except Exception as e:
            messagebox.showerror("Aktarma Hatası", f"Excel'e aktarılırken bir sorun oluştu:\n{str(e)}")

if __name__ == "__main__":
    root = tk.Tk()
    root.report_callback_exception = hata_yakalayici 
    app = DersProgramiApp(root)
    root.mainloop()