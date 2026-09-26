import traceback
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from ortools.sat.python import cp_model
import pandas as pd
import json
import os

# --- İPUCU (TOOLTIP) SINIFI ---
class ToolTip(object):
    def __init__(self, widget, text='widget info'):
        self.widget = widget
        self.text = text
        self.widget.bind("<Enter>", self.enter)
        self.widget.bind("<Leave>", self.close)
        self.tw = None

    def enter(self, event=None):
        x, y, cx, cy = self.widget.bbox("insert")
        x += self.widget.winfo_rootx() + 25
        y += self.widget.winfo_rooty() + 20
        self.tw = tk.Toplevel(self.widget)
        self.tw.wm_overrideredirect(True)
        self.tw.wm_geometry("+%d+%d" % (x, y))
        label = tk.Label(self.tw, text=self.text, justify='left',
                         background="#2c3e50", foreground="#ffffff", relief='flat', borderwidth=0,
                         font=("Segoe UI", 9, "normal"))
        label.pack(ipadx=6, ipady=4)

    def close(self, event=None):
        if self.tw:
            self.tw.destroy()
            self.tw = None

def hata_yakalayici(exc_type, exc_value, exc_traceback):
    hata_metni = "".join(traceback.format_exception(exc_type, exc_value, exc_traceback))
    hata_penceresi = tk.Toplevel()
    hata_penceresi.title("Beklenmeyen Sistem Hatası")
    hata_penceresi.geometry("700x400")
    hata_penceresi.attributes('-topmost', True)
    
    tk.Label(hata_penceresi, text="Programda arka planda bir hata oluştu. Lütfen aşağıdaki detayları inceleyin:", 
             font=("Segoe UI", 10, "bold"), fg="#D32F2F").pack(anchor="w", padx=10, pady=10)
    
    frame_text = tk.Frame(hata_penceresi)
    frame_text.pack(expand=True, fill="both", padx=10)
    scrollbar = ttk.Scrollbar(frame_text)
    scrollbar.pack(side="right", fill="y")
    metin_kutusu = tk.Text(frame_text, wrap="word", font=("Consolas", 9), yscrollcommand=scrollbar.set)
    metin_kutusu.insert(tk.END, hata_metni)
    metin_kutusu.config(state="disabled")
    metin_kutusu.pack(side="left", expand=True, fill="both")
    scrollbar.config(command=metin_kutusu.yview)
    
    tk.Button(hata_penceresi, text="Pencereyi Kapat", command=hata_penceresi.destroy, 
              bg="#ececec", font=("Segoe UI", 10)).pack(pady=10)

class DersProgramiApp:
    def __init__(self, root):
        self.root = root
        self.root.title("KMÜ Ders Programı Planlayıcısı")
        self.root.geometry("1150x920")
        
        # --- MODERN ARAYÜZ (UI) YAPILANDIRMASI ---
        self.style = ttk.Style()
        self.style.theme_use("clam")
        
        self.bg_color = "#F4F6F9"
        self.root.configure(bg=self.bg_color)
        
        self.font_standart = ("Segoe UI", 10)
        self.font_baslik = ("Segoe UI", 10, "bold")
        
        self.style.configure("TFrame", background=self.bg_color)
        self.style.configure("TLabel", background=self.bg_color, foreground="#2C3E50")
        self.style.configure("TCheckbutton", background=self.bg_color, font=self.font_standart)
        self.style.configure("TLabelframe", background=self.bg_color, borderwidth=1, bordercolor="#BDC3C7")
        self.style.configure("TLabelframe.Label", background=self.bg_color, foreground="#2980B9", font=("Segoe UI", 11, "bold"))
        
        self.style.configure("Treeview.Heading", font=self.font_baslik, background="#E0E6ED", foreground="#2C3E50", borderwidth=0)
        self.style.configure("Treeview", font=self.font_standart, rowheight=28, borderwidth=0)
        self.style.map("Treeview", background=[('selected', '#3498DB')], foreground=[('selected', 'white')])
        
        self.style.configure("TNotebook", background=self.bg_color, borderwidth=0)
        self.style.configure("TNotebook.Tab", padding=[15, 6], font=("Segoe UI", 10, "bold"), background="#E0E6ED", foreground="#7F8C8D")
        self.style.map("TNotebook.Tab", background=[("selected", "#3498DB")], foreground=[("selected", "white")])

        # AYRI DOSYA YAPILANDIRMASI
        self.veri_dosyasi = "kmu_veritabani.json"
        self.kilavuz_dosyasi = "kmu_kilavuz.json"
        
        self.bolum_adi = "BÖLÜMÜ"
        self.kilavuz_metni = ""
        self.ogretim_elemanlari = []
        self.siniflar = []
        self.dersler = []
        self.olusturulan_program = []

        self.gunler = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma"]
        self.saatler = [
            "08:30-09:15", "09:30-10:15", "10:30-11:15", "11:30-12:15",
            "13:00-13:45", "14:00-14:45", "15:00-15:45", "16:00-16:45"
        ]
        self.ogle_arasi_siniri = 4
        
        # MEKAN TİPLERİ LİSTESİ (Genişletildi)
        self.mekan_tipleri = [
            "Derslik (Genel)", 
            "Derslik Tip 1", 
            "Derslik Tip 2", 
            "Derslik Tip 3", 
            "Anfi Tip 1",
            "Anfi Tip 2",
            "Anfi Tip 3",
            "Genel Laboratuvar", 
            "Lab Tip 1", 
            "Lab Tip 2", 
            "Lab Tip 3", 
            "Lab Tip 4", 
            "Lab Tip 5"
        ]

        self.varsayilan_kilavuz = """=== KMÜ DERS PROGRAMI PLANLAYICISI KULLANIM KILAVUZU ===

Bu yazılım, akademik bölümlerin ders programı hazırlama süreçlerini optimize etmek amacıyla Kısıt Programlama (Constraint Programming) yaklaşımı temel alınarak geliştirilmiştir. Altyapısında Google'ın açık kaynaklı optimizasyon çözücüsü OR-Tools (Google Optimization Tools) ve Python programlama dili kullanılmıştır. Açık kaynak kodlu olarak lisanslanan bu sistem, Necati Buğra KUDDAŞ tarafından tasarlanmış ve geliştirilmiştir.

1. ÖĞRETİM ELEMANLARI SEKME KULLANIMI:
- Bölümde görev yapan tüm öğretim elemanlarını sisteme kaydedin.
- Sağ tarafta yer alan müsaitlik matrisi üzerinden, ilgili öğretim elemanının ders atanamayacak (müsait olmadığı) gün ve saatlerindeki tik işaretlerini kaldırın.

2. SINIFLAR VE LABORATUVARLAR SEKME KULLANIMI:
- Derslerin işleneceği fiziksel mekanları (Derslik, Anfi Tip 1, Lab Tip 2 vb.) sisteme tanımlayın. "Mekan Türü" seçimi, doğru eşleştirme için büyük önem taşımaktadır.
- Belirli bir mekanın başka bir programa tahsis edildiği veya kullanıma uygun olmadığı saatler varsa, ilgili saatlerdeki tik işaretlerini kaldırarak algoritmanın atama yapmasını engelleyin.

3. DERS HAVUZU SEKME KULLANIMI (KRİTİK AYARLAR):
- Derslik Gerektirmez: Staj, ofis uygulaması veya hastane uygulaması gibi fiziksel bir mekana ihtiyaç duymayan dersler için tercih edilmelidir.
- Dersin İşleneceği Mekan: Bu alandan dersin hangi tür laboratuvar, amfi veya derslikte işleneceğini seçebilirsiniz. Sistem, dersi kesinlikle seçilen türdeki bir mekana atayacaktır.
- Sabit Gün / Saat Kilidi: Kurum dışı etkenler veya zorunlu haller nedeniyle bir dersin kesin olarak belirli bir gün ve saatte işlenmesi gerekiyorsa, ilgili zaman dilimi bu alandan sabitlenmelidir.

4. PROGRAM ÜRETİMİ VE DIŞA AKTARIM:
- İlgili dönemi seçip "Programı Üret" butonuna tıkladığınızda, algoritma tüm kısıtlamaları ve on binlerce farklı kombinasyonu eşzamanlı olarak değerlendirerek çakışmasız en ideal programı oluşturur.
- "Çözüm Bulunamadı" veya "Darboğaz" şeklinde bir uyarı alırsanız, bu durum belirlenen kısıtlamaların matematiksel olarak bir çözüme izin vermediğini gösterir (Örn: Anfi Tip 1 isteyen 3 farklı ders varken sadece 1 adet Anfi Tip 1 mekanının olması gibi). Bu gibi durumlarda kısıtları esnetmeniz önerilir.
- Üretilen nihai ders programını, tek bir tıklama ile doğrudan kurumsal Excel formatında dışa aktarabilirsiniz.

(Not: Bu metin, program dizininde yer alan 'kmu_kilavuz.json' dosyası üzerinden dilediğiniz gibi güncellenebilir ve bölümünüze özgü hale getirilebilir.)
"""

        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=10)

        self.tab_ogr = ttk.Frame(self.notebook)
        self.tab_siniflar = ttk.Frame(self.notebook)
        self.tab_dersler = ttk.Frame(self.notebook)
        self.tab_sonuc = ttk.Frame(self.notebook)
        self.tab_kilavuz = ttk.Frame(self.notebook)

        self.notebook.add(self.tab_ogr, text="Öğretim Elemanları")
        self.notebook.add(self.tab_siniflar, text="Mekanlar & Laboratuvarlar")
        self.notebook.add(self.tab_dersler, text="Ders Havuzu")
        self.notebook.add(self.tab_sonuc, text="Program Üret & Aktar")
        self.notebook.add(self.tab_kilavuz, text="Kullanım Kılavuzu")

        self.ogr_musaitlik_vars = {}
        self.sinif_musaitlik_vars = {}
        self.derslik_gerekmez_var = tk.BooleanVar(value=False)
        self.aktif_mi_var = tk.BooleanVar(value=True)

        self.arayuz_ogretim_elemanlari()
        self.arayuz_siniflar()
        self.arayuz_dersler()
        self.arayuz_sonuc()
        self.arayuz_kilavuz()

        self.veri_yukle()
        self.kilavuz_yukle()

    def olustur_ipucu(self, ebeveyn, metin):
        lbl = tk.Label(ebeveyn, text="[?]", fg="#2980B9", bg=self.bg_color, cursor="hand2", font=("Segoe UI", 9, "bold"))
        ToolTip(lbl, metin)
        return lbl

    def modern_buton(self, parent, text, command, bg_color, fg_color="white", icon=""):
        display_text = f"{icon} {text}" if icon else text
        btn = tk.Button(parent, text=display_text, command=command, bg=bg_color, fg=fg_color, 
                        font=("Segoe UI", 10, "bold"), relief="flat", borderwidth=0, padx=12, pady=6, cursor="hand2")
        return btn

    def matris_buton(self, parent, text, command):
        btn = tk.Button(parent, text=text, font=("Segoe UI", 8, "bold"), relief="flat", bg="#D5D8DC", fg="#2C3E50", command=command, cursor="hand2")
        return btn

    def toggle_gun(self, gun, var_dict):
        ilk_durum = var_dict[(gun, self.saatler[0])].get()
        for saat in self.saatler: var_dict[(gun, saat)].set(not ilk_durum)

    def toggle_blok(self, is_sabah, var_dict):
        saatler_dilimi = self.saatler[:self.ogle_arasi_siniri] if is_sabah else self.saatler[self.ogle_arasi_siniri:]
        ilk_durum = var_dict[(self.gunler[0], saatler_dilimi[0])].get()
        for gun in self.gunler:
            for saat in saatler_dilimi: var_dict[(gun, saat)].set(not ilk_durum)

    def tumunu_sec(self, var_dict):
        ilk_durum = var_dict[(self.gunler[0], self.saatler[0])].get()
        for var in var_dict.values(): var.set(not ilk_durum)

    def form_sifirla_ogr(self):
        for item in self.tree_ogr.selection(): self.tree_ogr.selection_remove(item)
        self.entry_ogr_ad.delete(0, tk.END)
        for var in self.ogr_musaitlik_vars.values(): var.set(True)

    def form_sifirla_sinif(self):
        for item in self.tree_sinif.selection(): self.tree_sinif.selection_remove(item)
        self.entry_sinif.delete(0, tk.END)
        self.combo_sinif_tipi.set("Derslik (Genel)")
        for var in self.sinif_musaitlik_vars.values(): var.set(True)

    def form_sifirla_ders(self):
        for tree in [self.tree_2026, self.tree_2024, self.tree_diger]:
            for item in tree.selection(): tree.selection_remove(item)
        self.entry_ders.delete(0, tk.END)
        self.entry_sure.delete(0, tk.END)
        self.listbox_ogr.selection_clear(0, tk.END)
        self.derslik_gerekmez_var.set(False)
        self.combo_istenen_mekan.set("Derslik (Genel)")
        self.aktif_mi_var.set(True)
        self.combo_sabit_gun.set("(Esnek / Serbest)")
        self.combo_sabit_saat.set("(Esnek / Serbest)")

    # --- 1. SEKME ---
    def arayuz_ogretim_elemanlari(self):
        frame_bilgi = ttk.Frame(self.tab_ogr)
        frame_bilgi.pack(fill="x", padx=10, pady=(10, 0))
        ttk.Label(frame_bilgi, text="📌 Bölümdeki öğretim elemanlarını sisteme kaydedin. Müsait olunmayan saatlerin tik işaretlerini kaldırarak sistemi sınırlandırın.", font=("Segoe UI", 10, "italic"), foreground="#7F8C8D").pack(anchor="w")

        frame_sol = ttk.Frame(self.tab_ogr)
        frame_sol.pack(side="left", fill="y", padx=10, pady=10)
        frame_sag = ttk.Frame(self.tab_ogr)
        frame_sag.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        frame_ust = ttk.Frame(frame_sol)
        frame_ust.pack(fill="x", pady=5)
        ttk.Label(frame_ust, text="Öğretim Elemanı Adı:").pack(side="left")
        self.entry_ogr_ad = ttk.Entry(frame_ust, width=25, font=self.font_standart)
        self.entry_ogr_ad.pack(side="right", fill="x", expand=True, padx=5)

        frame_matris = ttk.LabelFrame(frame_sol, padding=10)
        frame_matris.pack(fill="x", pady=10)
        
        frame_matris_baslik = ttk.Frame(frame_matris)
        frame_matris_baslik.grid(row=0, column=0, columnspan=6, sticky="w", pady=(0,5))
        ttk.Label(frame_matris_baslik, text="Müsaitlik Matrisi", font=self.font_baslik).pack(side="left")
        self.olustur_ipucu(frame_matris_baslik, "Tik işaretleri kaldırılan saatlere, algoritma tarafından kesinlikle ders atanamaz.").pack(side="left", padx=5)

        for j, gun in enumerate(self.gunler):
            self.matris_buton(frame_matris, gun, lambda g=gun: self.toggle_gun(g, self.ogr_musaitlik_vars)).grid(row=1, column=j+1, padx=2, pady=2, sticky="we")
            
        for i, saat in enumerate(self.saatler):
            row_idx = i + 2
            if i >= self.ogle_arasi_siniri: row_idx += 1
            if i == self.ogle_arasi_siniri:
                ttk.Label(frame_matris, text="--- ÖĞLE ARASI ---", foreground="#E74C3C", font=("Segoe UI", 8, "bold"), anchor="center").grid(row=row_idx-1, column=0, columnspan=6, pady=2)
            ttk.Label(frame_matris, text=saat, font=("Segoe UI", 8)).grid(row=row_idx, column=0, sticky="e", padx=5)
            for j, gun in enumerate(self.gunler):
                var = tk.BooleanVar(value=True)
                self.ogr_musaitlik_vars[(gun, saat)] = var
                ttk.Checkbutton(frame_matris, variable=var).grid(row=row_idx, column=j+1)

        frame_hizli = ttk.Frame(frame_matris)
        frame_hizli.grid(row=len(self.saatler)+3, column=0, columnspan=6, pady=10)
        self.matris_buton(frame_hizli, "Sabah", lambda: self.toggle_blok(True, self.ogr_musaitlik_vars)).pack(side="left", padx=2)
        self.matris_buton(frame_hizli, "Öğleden Sonra", lambda: self.toggle_blok(False, self.ogr_musaitlik_vars)).pack(side="left", padx=2)
        self.matris_buton(frame_hizli, "Tümünü Seç/Kaldır", lambda: self.tumunu_sec(self.ogr_musaitlik_vars)).pack(side="left", padx=2)

        frame_liste_buton = ttk.Frame(frame_sol)
        frame_liste_buton.pack(fill="x", pady=5)
        self.modern_buton(frame_liste_buton, "Sisteme Ekle / Güncelle", self.ogr_ekle, "#27AE60").pack(fill="x", pady=3)
        self.modern_buton(frame_liste_buton, "Seçimi Temizle (Yeni Kayıt)", self.form_sifirla_ogr, "#F39C12").pack(fill="x", pady=3)

        ttk.Label(frame_sag, text="Kayıtlı Öğretim Elemanları", font=self.font_baslik).pack(anchor="w", pady=(0,5))
        self.tree_ogr = ttk.Treeview(frame_sag, columns=("Ad"), show="headings")
        self.tree_ogr.heading("Ad", text="Öğretim Elemanı")
        self.tree_ogr.pack(fill="both", expand=True, pady=5)
        self.tree_ogr.bind('<<TreeviewSelect>>', self.ogr_secildi)
        
        self.modern_buton(frame_sag, "Seçili Olanı Sil", self.ogr_sil, "#E74C3C").pack(anchor="e", pady=5)

    # --- 2. SEKME ---
    def arayuz_siniflar(self):
        frame_bilgi = ttk.Frame(self.tab_siniflar)
        frame_bilgi.pack(fill="x", padx=10, pady=(10, 0))
        ttk.Label(frame_bilgi, text="📌 Derslik ve Laboratuvarları ekleyin. Mekanın kullanıma uygun olmadığı zamanlar varsa tik işaretlerini kaldırın.", font=("Segoe UI", 10, "italic"), foreground="#7F8C8D").pack(anchor="w")

        frame_sol = ttk.Frame(self.tab_siniflar)
        frame_sol.pack(side="left", fill="y", padx=10, pady=10)
        frame_sag = ttk.Frame(self.tab_siniflar)
        frame_sag.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        frame_ust = ttk.Frame(frame_sol)
        frame_ust.pack(fill="x", pady=5)
        ttk.Label(frame_ust, text="Mekan Adı:").pack(side="left")
        self.entry_sinif = ttk.Entry(frame_ust, width=20, font=self.font_standart)
        self.entry_sinif.pack(side="right", fill="x", expand=True, padx=5)

        frame_tip = ttk.Frame(frame_sol)
        frame_tip.pack(fill="x", pady=5)
        ttk.Label(frame_tip, text="Mekan Türü:").pack(side="left")
        self.olustur_ipucu(frame_tip, "Seçilen türe göre, buraya yalnızca aynı kategoriye (Derslik, Anfi, Lab vb.) sahip dersler atanacaktır.").pack(side="left", padx=5)
        self.combo_sinif_tipi = ttk.Combobox(frame_tip, values=self.mekan_tipleri, width=17, state="readonly", font=self.font_standart)
        self.combo_sinif_tipi.set("Derslik (Genel)")
        self.combo_sinif_tipi.pack(side="right", padx=5)

        frame_matris = ttk.LabelFrame(frame_sol, padding=10)
        frame_matris.pack(fill="x", pady=10)

        frame_matris_baslik = ttk.Frame(frame_matris)
        frame_matris_baslik.grid(row=0, column=0, columnspan=6, sticky="w", pady=(0,5))
        ttk.Label(frame_matris_baslik, text="Müsaitlik Matrisi", font=self.font_baslik).pack(side="left")

        for j, gun in enumerate(self.gunler):
            self.matris_buton(frame_matris, gun, lambda g=gun: self.toggle_gun(g, self.sinif_musaitlik_vars)).grid(row=1, column=j+1, padx=2, pady=2, sticky="we")
            
        for i, saat in enumerate(self.saatler):
            row_idx = i + 2
            if i >= self.ogle_arasi_siniri: row_idx += 1
            if i == self.ogle_arasi_siniri:
                ttk.Label(frame_matris, text="--- ÖĞLE ARASI ---", foreground="#E74C3C", font=("Segoe UI", 8, "bold"), anchor="center").grid(row=row_idx-1, column=0, columnspan=6, pady=2)
            ttk.Label(frame_matris, text=saat, font=("Segoe UI", 8)).grid(row=row_idx, column=0, sticky="e", padx=5)
            for j, gun in enumerate(self.gunler):
                var = tk.BooleanVar(value=True)
                self.sinif_musaitlik_vars[(gun, saat)] = var
                ttk.Checkbutton(frame_matris, variable=var).grid(row=row_idx, column=j+1)

        frame_hizli = ttk.Frame(frame_matris)
        frame_hizli.grid(row=len(self.saatler)+3, column=0, columnspan=6, pady=10)
        self.matris_buton(frame_hizli, "Sabah", lambda: self.toggle_blok(True, self.sinif_musaitlik_vars)).pack(side="left", padx=2)
        self.matris_buton(frame_hizli, "Öğleden Sonra", lambda: self.toggle_blok(False, self.sinif_musaitlik_vars)).pack(side="left", padx=2)
        self.matris_buton(frame_hizli, "Tümünü Seç/Kaldır", lambda: self.tumunu_sec(self.sinif_musaitlik_vars)).pack(side="left", padx=2)

        frame_liste_buton = ttk.Frame(frame_sol)
        frame_liste_buton.pack(fill="x", pady=5)
        self.modern_buton(frame_liste_buton, "Sisteme Ekle / Güncelle", self.sinif_ekle, "#27AE60").pack(fill="x", pady=3)
        self.modern_buton(frame_liste_buton, "Seçimi Temizle (Yeni Kayıt)", self.form_sifirla_sinif, "#F39C12").pack(fill="x", pady=3)

        ttk.Label(frame_sag, text="Kayıtlı Sınıf ve Laboratuvarlar", font=self.font_baslik).pack(anchor="w", pady=(0,5))
        self.tree_sinif = ttk.Treeview(frame_sag, columns=("Ad", "Tip"), show="headings")
        self.tree_sinif.heading("Ad", text="Mekan Adı")
        self.tree_sinif.heading("Tip", text="Tür")
        self.tree_sinif.column("Ad", width=180)
        self.tree_sinif.column("Tip", width=100, anchor="center")
        self.tree_sinif.pack(fill="both", expand=True, pady=5)
        self.tree_sinif.bind('<<TreeviewSelect>>', self.sinif_secildi)
        
        self.modern_buton(frame_sag, "Seçili Olanı Sil", self.sinif_sil, "#E74C3C").pack(anchor="e", pady=5)

    # --- 3. SEKME ---
    def arayuz_dersler(self):
        frame_bilgi = ttk.Frame(self.tab_dersler)
        frame_bilgi.pack(fill="x", padx=10, pady=(10, 0))
        ttk.Label(frame_bilgi, text="📌 Ders havuzunu oluşturun. Mekan tipi atamalarını, mekansız uygulamaları veya sabit gün/saat kilitlerini buradan yönetebilirsiniz.", font=("Segoe UI", 10, "italic"), foreground="#7F8C8D").pack(anchor="w")

        frame_sol = ttk.Frame(self.tab_dersler)
        frame_sol.pack(side="left", fill="y", padx=10, pady=10)
        frame_sag = ttk.Frame(self.tab_dersler)
        frame_sag.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        frame_giris = ttk.LabelFrame(frame_sol, text="Ders ve Müfredat Bilgileri", padding=10)
        frame_giris.pack(fill="x")

        ttk.Label(frame_giris, text="Dersin Adı:").grid(row=0, column=0, sticky="w", pady=5)
        self.entry_ders = ttk.Entry(frame_giris, width=30, font=self.font_standart)
        self.entry_ders.grid(row=0, column=1, padx=5, pady=5)

        ttk.Label(frame_giris, text="Öğr. Elemanı:\n(Birden fazla seçilebilir)").grid(row=1, column=0, sticky="nw", pady=5)
        frame_listbox_ogr = ttk.Frame(frame_giris)
        frame_listbox_ogr.grid(row=1, column=1, padx=5, pady=5, sticky="we")
        
        self.listbox_ogr = tk.Listbox(frame_listbox_ogr, selectmode=tk.MULTIPLE, height=4, font=self.font_standart, exportselection=False, relief="flat", highlightthickness=1, highlightcolor="#BDC3C7")
        scrollbar_ogr = ttk.Scrollbar(frame_listbox_ogr, orient="vertical", command=self.listbox_ogr.yview)
        self.listbox_ogr.configure(yscrollcommand=scrollbar_ogr.set)
        self.listbox_ogr.pack(side="left", fill="both", expand=True)
        scrollbar_ogr.pack(side="right", fill="y")

        frame_sure_lbl = ttk.Frame(frame_giris)
        frame_sure_lbl.grid(row=2, column=0, sticky="nw", pady=5)
        ttk.Label(frame_sure_lbl, text="Süre (Blok):").pack(side="left")
        self.olustur_ipucu(frame_sure_lbl, "Örn: 3 yazarsanız, ders 3 saat boyunca bölünmeden işlenir.\n8 yazılması durumunda sistem bunu tam gün uygulaması olarak değerlendirir.").pack(side="left", padx=2)
        
        self.entry_sure = ttk.Entry(frame_giris, width=30, font=self.font_standart)
        self.entry_sure.grid(row=2, column=1, padx=5, pady=5)

        ttk.Label(frame_giris, text="Sınıf Seviyesi:").grid(row=3, column=0, sticky="w", pady=5)
        self.combo_sinif_yili = ttk.Combobox(frame_giris, values=["1. Sınıf", "2. Sınıf", "3. Sınıf", "4. Sınıf", "Seçmeli/Karma"], width=28, state="readonly", font=self.font_standart)
        self.combo_sinif_yili.set("1. Sınıf")
        self.combo_sinif_yili.grid(row=3, column=1, padx=5, pady=5)

        ttk.Label(frame_giris, text="Dönem / Yıl:").grid(row=4, column=0, sticky="w", pady=5)
        frame_donem = ttk.Frame(frame_giris)
        frame_donem.grid(row=4, column=1, padx=5, pady=5, sticky="we")
        self.combo_donem_kayit = ttk.Combobox(frame_donem, values=["Güz", "Bahar", "Yaz"], width=10, state="readonly", font=self.font_standart)
        self.combo_donem_kayit.set("Güz")
        self.combo_donem_kayit.pack(side="left", padx=(0,5))
        self.combo_mufredat = ttk.Combobox(frame_donem, values=["2023", "2024", "2026", "Ortak"], width=10, state="readonly", font=self.font_standart)
        self.combo_mufredat.set("2026")
        self.combo_mufredat.pack(side="left")

        # ÖZEL AYARLAR
        frame_kurallar_dis = ttk.LabelFrame(frame_sol, padding=10)
        frame_kurallar_dis.pack(fill="x", pady=5)
        
        frame_kural_baslik = ttk.Frame(frame_kurallar_dis)
        frame_kural_baslik.grid(row=0, column=0, columnspan=2, sticky="w", pady=(0,5))
        ttk.Label(frame_kural_baslik, text="Özel Ayarlar", font=self.font_baslik).pack(side="left")
        self.olustur_ipucu(frame_kural_baslik, "Pasif konuma getirilen dersler, program optimizasyonu aşamasında dikkate alınmaz.\nDerslik gerektirmez ibaresi seçildiğinde ders herhangi bir mekana atanmaz.").pack(side="left", padx=5)

        ttk.Checkbutton(frame_kurallar_dis, text="Ders Bu Dönem AÇILACAK (Aktif)", variable=self.aktif_mi_var).grid(row=1, column=0, columnspan=2, sticky="w", pady=2)
        ttk.Checkbutton(frame_kurallar_dis, text="Derslik Gerektirmez (Ofis/Dış Uygulama)", variable=self.derslik_gerekmez_var).grid(row=2, column=0, columnspan=2, sticky="w", pady=2)

        frame_mekan_tip = ttk.Frame(frame_kurallar_dis)
        frame_mekan_tip.grid(row=3, column=0, columnspan=2, sticky="w", pady=(5, 2))
        ttk.Label(frame_mekan_tip, text="Dersin İşleneceği Mekan:").pack(side="left")
        self.combo_istenen_mekan = ttk.Combobox(frame_mekan_tip, values=self.mekan_tipleri, width=17, state="readonly", font=self.font_standart)
        self.combo_istenen_mekan.set("Derslik (Genel)")
        self.combo_istenen_mekan.pack(side="left", padx=5)

        # SABİT GÜN / SAAT KİLİDİ
        frame_sabit = ttk.LabelFrame(frame_sol, padding=10)
        frame_sabit.pack(fill="x", pady=5)

        frame_sabit_baslik = ttk.Frame(frame_sabit)
        frame_sabit_baslik.grid(row=0, column=0, columnspan=2, sticky="w", pady=(0,5))
        ttk.Label(frame_sabit_baslik, text="Sabit Gün / Saat Kilidi", font=self.font_baslik).pack(side="left")
        self.olustur_ipucu(frame_sabit_baslik, "Dersin belirtilen zaman diliminde sabitlenmesini sağlar.\nAlgoritma bu dersin gün ve saatini kesinlikle değiştiremez.").pack(side="left", padx=5)

        ttk.Label(frame_sabit, text="Sabit Gün:").grid(row=1, column=0, sticky="w", pady=2)
        self.combo_sabit_gun = ttk.Combobox(frame_sabit, values=["(Esnek / Serbest)"] + self.gunler, width=18, state="readonly", font=self.font_standart)
        self.combo_sabit_gun.set("(Esnek / Serbest)")
        self.combo_sabit_gun.grid(row=1, column=1, sticky="w", pady=2, padx=5, columnspan=2)

        ttk.Label(frame_sabit, text="Sabit Saat:").grid(row=2, column=0, sticky="w", pady=2)
        self.combo_sabit_saat = ttk.Combobox(frame_sabit, values=["(Esnek / Serbest)"] + self.saatler, width=18, state="readonly", font=self.font_standart)
        self.combo_sabit_saat.set("(Esnek / Serbest)")
        self.combo_sabit_saat.grid(row=2, column=1, sticky="w", pady=2, padx=5, columnspan=2)

        frame_liste_buton = ttk.Frame(frame_sol)
        frame_liste_buton.pack(fill="x", pady=10)
        self.modern_buton(frame_liste_buton, "Havuza Ekle / Güncelle", self.ders_ekle, "#27AE60").pack(fill="x", pady=3)
        self.modern_buton(frame_liste_buton, "Seçimi Temizle (Yeni Kayıt)", self.form_sifirla_ders, "#F39C12").pack(fill="x", pady=3)

        ttk.Label(frame_sag, text="Veritabanındaki Dersler", font=self.font_baslik).pack(anchor="w", pady=(0, 5))
        
        self.notebook_mufredat = ttk.Notebook(frame_sag)
        self.notebook_mufredat.pack(fill="both", expand=True)
        self.notebook_mufredat.bind("<<NotebookTabChanged>>", lambda e: self.form_sifirla_ders())

        frame_2026 = ttk.Frame(self.notebook_mufredat)
        frame_2024 = ttk.Frame(self.notebook_mufredat)
        frame_diger = ttk.Frame(self.notebook_mufredat)

        self.notebook_mufredat.add(frame_2026, text="   2026 Müfredatı   ")
        self.notebook_mufredat.add(frame_2024, text="   2024 Müfredatı   ")
        self.notebook_mufredat.add(frame_diger, text="   Diğer (2023/Ortak)   ")

        self.tree_2026 = self.ders_tablosu_olustur(frame_2026)
        self.tree_2024 = self.ders_tablosu_olustur(frame_2024)
        self.tree_diger = self.ders_tablosu_olustur(frame_diger)

        frame_sag_alt = ttk.Frame(frame_sag)
        frame_sag_alt.pack(fill="x", pady=5)
        self.modern_buton(frame_sag_alt, "Seçili Dersleri AKTİF YAP", lambda: self.ders_toplu_durum(True), "#3498DB").pack(side="left", padx=2)
        self.modern_buton(frame_sag_alt, "Seçili Dersleri PASİF YAP", lambda: self.ders_toplu_durum(False), "#95A5A6").pack(side="left", padx=2)
        self.modern_buton(frame_sag_alt, "Seçili Dersi SİL", self.ders_sil, "#E74C3C").pack(side="right", padx=2)

    def ders_tablosu_olustur(self, parent):
        tree = ttk.Treeview(parent, columns=("Durum", "Donem", "Ders", "Sinif", "Ogr", "Sure", "Mekan", "Sabit"), show="headings", selectmode="extended")
        tree.heading("Durum", text="Durum")
        tree.heading("Donem", text="Dönem")
        tree.heading("Ders", text="Ders Adı")
        tree.heading("Sinif", text="Seviye")
        tree.heading("Ogr", text="Öğr. Elemanı")
        tree.heading("Sure", text="Süre")
        tree.heading("Mekan", text="Mekan/Tip")
        tree.heading("Sabit", text="Sabit Kilit")
        
        tree.column("Durum", width=55, anchor="center")
        tree.column("Donem", width=45, anchor="center")
        tree.column("Ders", width=140)
        tree.column("Sinif", width=65, anchor="center")
        tree.column("Ogr", width=120)
        tree.column("Sure", width=35, anchor="center")
        tree.column("Mekan", width=95, anchor="center")
        tree.column("Sabit", width=85, anchor="center")
        
        tree.pack(fill="both", expand=True, padx=2, pady=2)
        tree.bind('<<TreeviewSelect>>', self.ders_secildi)
        return tree

    def aktif_tabloyu_getir(self):
        tab_id = self.notebook_mufredat.index("current")
        if tab_id == 0: return self.tree_2026
        elif tab_id == 1: return self.tree_2024
        else: return self.tree_diger

    # --- 4. SEKME ---
    def arayuz_sonuc(self):
        frame_bilgi = ttk.Frame(self.tab_sonuc)
        frame_bilgi.pack(fill="x", padx=10, pady=(10, 0))
        ttk.Label(frame_bilgi, text="📌 Tüm kısıtlamalara göre çakışmasız programınızı üretin. Başarılı sonuç elde edildiğinde Excel tablosuna aktarabilirsiniz.", font=("Segoe UI", 10, "italic"), foreground="#7F8C8D").pack(anchor="w")

        frame_ust = ttk.Frame(self.tab_sonuc)
        frame_ust.pack(fill="x", padx=10, pady=10)

        frame_donem = ttk.LabelFrame(frame_ust, text="Program Optimizasyonu", padding=15)
        frame_donem.pack(fill="x", expand=True)

        ttk.Label(frame_donem, text="Program Üretilecek Dönem:").pack(side="left", padx=5)
        self.combo_donem_secim = ttk.Combobox(frame_donem, values=["Güz", "Bahar", "Yaz"], state="readonly", width=15, font=self.font_standart)
        self.combo_donem_secim.set("Güz")
        self.combo_donem_secim.pack(side="left", padx=10)

        frame_buton = ttk.Frame(frame_donem)
        frame_buton.pack(side="left", padx=20)

        self.modern_buton(frame_buton, "Yapay Zeka İle Programı Üret", self.program_uret, "#2980B9").pack(side="left")
        self.olustur_ipucu(frame_buton, "Algoritma tüm kısıtlamaları matematiksel olarak analiz ederek\nçakışmasız en ideal kombinasyonu bulur.").pack(side="left", padx=5)

        self.modern_buton(frame_donem, "Excel'e Aktar", self.excel_aktar, "#27AE60").pack(side="right", padx=10)

        self.tree_sonuc = ttk.Treeview(self.tab_sonuc, columns=("Gun", "Saat", "Derslik", "Ders", "Sinif", "Ogr"), show="headings")
        self.tree_sonuc.heading("Gun", text="Gün")
        self.tree_sonuc.heading("Saat", text="Saat")
        self.tree_sonuc.heading("Derslik", text="Derslik/Mekan")
        self.tree_sonuc.heading("Ders", text="Ders Adı")
        self.tree_sonuc.heading("Sinif", text="Seviye")
        self.tree_sonuc.heading("Ogr", text="Öğr. Elemanı")
        
        self.tree_sonuc.column("Gun", width=90, anchor="center")
        self.tree_sonuc.column("Saat", width=90, anchor="center")
        self.tree_sonuc.column("Derslik", width=120, anchor="center")
        self.tree_sonuc.column("Ders", width=200)
        self.tree_sonuc.column("Sinif", width=80, anchor="center")
        self.tree_sonuc.column("Ogr", width=180)

        self.tree_sonuc.pack(padx=10, pady=10, fill="both", expand=True)

    # --- 5. SEKME (KILAVUZ) ---
    def arayuz_kilavuz(self):
        frame_bilgi = ttk.Frame(self.tab_kilavuz)
        frame_bilgi.pack(fill="x", padx=10, pady=(10, 0))
        ttk.Label(frame_bilgi, text="📌 Bu kılavuz, yazılımın akademik standartlarını ve kullanım yönergelerini içerir.", font=("Segoe UI", 10, "italic"), foreground="#7F8C8D").pack(anchor="w")

        frame_metin = ttk.Frame(self.tab_kilavuz)
        frame_metin.pack(fill="both", expand=True, padx=10, pady=10)

        scrollbar = ttk.Scrollbar(frame_metin)
        scrollbar.pack(side="right", fill="y")

        self.text_kilavuz = tk.Text(frame_metin, wrap="word", font=("Segoe UI", 11), yscrollcommand=scrollbar.set, bg="#ffffff", fg="#2C3E50", padx=15, pady=15, relief="flat")
        self.text_kilavuz.pack(side="left", fill="both", expand=True)
        scrollbar.config(command=self.text_kilavuz.yview)

    # --- LİSTE VE TIKLAMA OLAYLARI ---
    def ogr_secildi(self, event):
        secili = self.tree_ogr.selection()
        if not secili: return
        idx = self.tree_ogr.index(secili[0])
        o = self.ogretim_elemanlari[idx]
        self.entry_ogr_ad.delete(0, tk.END)
        self.entry_ogr_ad.insert(0, o["ad"])
        for (g, s), v in self.ogr_musaitlik_vars.items():
            v.set(o["musaitlik"].get((g, s), True))

    def sinif_secildi(self, event):
        secili = self.tree_sinif.selection()
        if not secili: return
        idx = self.tree_sinif.index(secili[0])
        s = self.siniflar[idx]
        self.entry_sinif.delete(0, tk.END)
        self.entry_sinif.insert(0, s["ad"])
        
        # Geçmişe dönük uyumluluk dönüşümü
        s_tip = s.get("tip", "Derslik (Genel)")
        if s_tip == "Derslik": s_tip = "Derslik (Genel)"
        
        self.combo_sinif_tipi.set(s_tip)
        for (g, st), v in self.sinif_musaitlik_vars.items():
            v.set(s["musaitlik"].get((g, st), True))

    def ders_secildi(self, event):
        aktif_agac = self.aktif_tabloyu_getir()
        secili = aktif_agac.selection()
        if not secili: return
        
        idx = int(aktif_agac.item(secili[0], "tags")[0])
        d = self.dersler[idx]
        
        self.entry_ders.delete(0, tk.END)
        self.entry_ders.insert(0, d["id"])
        
        self.listbox_ogr.selection_clear(0, tk.END)
        for i, h in enumerate(self.ogretim_elemanlari):
            if h["ad"] in d["hocalar"]: self.listbox_ogr.selection_set(i)
                
        self.entry_sure.delete(0, tk.END)
        self.entry_sure.insert(0, str(d["sure"]))
        self.combo_sinif_yili.set(d.get("sinif_yil", "1. Sınıf"))
        self.combo_mufredat.set(d.get("mufredat", "2026"))
        self.combo_donem_kayit.set(d.get("donem", "Güz"))
        self.derslik_gerekmez_var.set(d.get("derslik_gerekmez", False))
        
        istenen = d.get("istenen_mekan", "Derslik (Genel)")
        if istenen == "Derslik": istenen = "Derslik (Genel)"
        if d.get("lab_gerekir", False) and "istenen_mekan" not in d:
            istenen = "Genel Laboratuvar"
        self.combo_istenen_mekan.set(istenen)
        
        self.aktif_mi_var.set(d.get("aktif_mi", True))
        self.combo_sabit_gun.set(d.get("sabit_gun", "(Esnek / Serbest)"))
        self.combo_sabit_saat.set(d.get("sabit_saat", "(Esnek / Serbest)"))

    # --- VERİ İŞLEMLERİ ---
    def kilavuz_yukle(self):
        guncelleme_gerekli = False
        if os.path.exists(self.kilavuz_dosyasi):
            try:
                with open(self.kilavuz_dosyasi, "r", encoding="utf-8") as f:
                    veri = json.load(f)
                    self.kilavuz_metni = veri.get("kilavuz_metni", self.varsayilan_kilavuz)
                    if "Necati Buğra KUDDAŞ" not in self.kilavuz_metni:
                        guncelleme_gerekli = True
            except:
                guncelleme_gerekli = True
        else:
            guncelleme_gerekli = True

        if guncelleme_gerekli:
            self.kilavuz_metni = self.varsayilan_kilavuz
            try:
                with open(self.kilavuz_dosyasi, "w", encoding="utf-8") as f:
                    json.dump({"kilavuz_metni": self.kilavuz_metni}, f, ensure_ascii=False, indent=4)
            except:
                pass
                
        self.text_kilavuz.config(state="normal")
        self.text_kilavuz.delete("1.0", tk.END)
        self.text_kilavuz.insert(tk.END, self.kilavuz_metni)
        self.text_kilavuz.config(state="disabled")

    def veri_kaydet(self):
        kayit_verisi = {
            "bolum_adi": self.bolum_adi,
            "ogretim_elemanlari": [
                {"ad": o["ad"], "musaitlik": {f"{g}_{s}": v for (g, s), v in o["musaitlik"].items()}} for o in self.ogretim_elemanlari
            ],
            "siniflar": [
                {"ad": s["ad"], "tip": s.get("tip", "Derslik (Genel)"), "musaitlik": {f"{g}_{s}": v for (g, s), v in s["musaitlik"].items()}} for s in self.siniflar
            ],
            "dersler": self.dersler
        }
        with open(self.veri_dosyasi, "w", encoding="utf-8") as f:
            json.dump(kayit_verisi, f, ensure_ascii=False, indent=4)

    def veri_yukle(self):
        if not os.path.exists(self.veri_dosyasi): return
        try:
            with open(self.veri_dosyasi, "r", encoding="utf-8") as f:
                kayit_verisi = json.load(f)

            self.bolum_adi = kayit_verisi.get("bolum_adi", "BÖLÜMÜ")
            self.root.title(f"KMÜ {self.bolum_adi} Ders Programı Planlayıcısı")

            for o in kayit_verisi.get("ogretim_elemanlari", []):
                musaitlik = {tuple(k.split("_")): v for k, v in o["musaitlik"].items()}
                self.ogretim_elemanlari.append({"ad": o["ad"], "musaitlik": musaitlik})
                self.tree_ogr.insert("", tk.END, values=(o["ad"],))

            for s in kayit_verisi.get("siniflar", []):
                musaitlik = {tuple(k.split("_")): v for k, v in s["musaitlik"].items()}
                s_tip = s.get("tip", "Derslik (Genel)")
                if s_tip == "Derslik": s_tip = "Derslik (Genel)"
                if s_tip == "Laboratuvar": s_tip = "Genel Laboratuvar"
                
                self.siniflar.append({"ad": s["ad"], "tip": s_tip, "musaitlik": musaitlik})
                self.tree_sinif.insert("", tk.END, values=(s["ad"], s_tip))

            for d in kayit_verisi.get("dersler", []):
                hoca_listesi = d.get("hocalar", d.get("hoca", ["Atanmadı"]))
                if isinstance(hoca_listesi, str): hoca_listesi = [hoca_listesi]
                
                istenen_mekan = d.get("istenen_mekan", "Derslik (Genel)")
                if istenen_mekan == "Derslik": istenen_mekan = "Derslik (Genel)"
                if d.get("lab_gerekir", False) and "istenen_mekan" not in d:
                    istenen_mekan = "Genel Laboratuvar"
                
                d_dict = {
                    "id": d["id"], "hocalar": hoca_listesi, "sure": d["sure"],
                    "sinif_yil": d.get("sinif_yil", "Belirtilmemiş"), "mufredat": d.get("mufredat", "2026"), 
                    "donem": d.get("donem", "Güz"), "derslik_gerekmez": d.get("derslik_gerekmez", False),
                    "istenen_mekan": istenen_mekan, "aktif_mi": d.get("aktif_mi", True),
                    "sabit_gun": d.get("sabit_gun", "(Esnek / Serbest)"), "sabit_saat": d.get("sabit_saat", "(Esnek / Serbest)")
                }
                self.dersler.append(d_dict)

            self.listbox_guncelle()
            self.ders_tablolarini_yenile()

        except Exception as e:
            messagebox.showerror("Hata", f"Veriler yüklenemedi:\n{str(e)}")

    def listbox_guncelle(self):
        self.listbox_ogr.delete(0, tk.END)
        for o in self.ogretim_elemanlari:
            self.listbox_ogr.insert(tk.END, o["ad"])

    def ders_tablolarini_yenile(self):
        for tree in [self.tree_2026, self.tree_2024, self.tree_diger]:
            for item in tree.get_children(): tree.delete(item)

        for i, d in enumerate(self.dersler):
            durum_metni = "✅ Aktif" if d["aktif_mi"] else "❌ Pasif"
            hoca_metni = ", ".join(d["hocalar"])
            
            if d.get("derslik_gerekmez", False):
                mekan_metni = "Ofis/Dış"
            else:
                mekan_metni = d.get("istenen_mekan", "Derslik (Genel)")
            
            sabit_gun_str = d.get('sabit_gun', '(Esnek / Serbest)')
            sabit_saat_str = d.get('sabit_saat', '(Esnek / Serbest)')
            
            if sabit_gun_str != "(Esnek / Serbest)" and sabit_saat_str != "(Esnek / Serbest)":
                sabit_metni = f"{sabit_gun_str[:3]} {sabit_saat_str[:5]}"
            elif sabit_gun_str != "(Esnek / Serbest)":
                sabit_metni = f"{sabit_gun_str[:3]} (Serb. Saat)"
            elif sabit_saat_str != "(Esnek / Serbest)":
                sabit_metni = f"Her Gün {sabit_saat_str[:5]}"
            else:
                sabit_metni = "Serbest"
            
            degerler = (durum_metni, d["donem"], d["id"], d["sinif_yil"], hoca_metni, d["sure"], mekan_metni, sabit_metni)
            
            if d["mufredat"] == "2026":
                self.tree_2026.insert("", tk.END, values=degerler, tags=(str(i),))
            elif d["mufredat"] == "2024":
                self.tree_2024.insert("", tk.END, values=degerler, tags=(str(i),))
            else:
                self.tree_diger.insert("", tk.END, values=degerler, tags=(str(i),))

    def ogr_ekle(self):
        ad = self.entry_ogr_ad.get().strip()
        if not ad: return messagebox.showwarning("Uyarı", "Ad boş olamaz.")
        musaitlik_durumu = {(g, s): self.ogr_musaitlik_vars[(g, s)].get() for g in self.gunler for s in self.saatler}
        secili = self.tree_ogr.selection()
        if secili:
            idx = self.tree_ogr.index(secili[0])
            self.ogretim_elemanlari[idx] = {"ad": ad, "musaitlik": musaitlik_durumu}
            self.tree_ogr.item(secili[0], values=(ad,))
            self.listbox_guncelle()
        else:
            if any(o["ad"] == ad for o in self.ogretim_elemanlari): return messagebox.showwarning("Uyarı", "Bu isim zaten var.")
            self.ogretim_elemanlari.append({"ad": ad, "musaitlik": musaitlik_durumu})
            self.tree_ogr.insert("", tk.END, values=(ad,))
            self.listbox_guncelle()
        self.veri_kaydet()

    def ogr_sil(self):
        secili = self.tree_ogr.selection()
        if not secili: return
        idx = self.tree_ogr.index(secili[0])
        del self.ogretim_elemanlari[idx]
        self.tree_ogr.delete(secili[0])
        self.listbox_guncelle()
        self.veri_kaydet()
        self.form_sifirla_ogr()

    def sinif_ekle(self):
        ad = self.entry_sinif.get().strip()
        s_tip = self.combo_sinif_tipi.get()
        if not ad: return messagebox.showwarning("Uyarı", "Mekan adı boş olamaz.")
        musaitlik_durumu = {(g, s): self.sinif_musaitlik_vars[(g, s)].get() for g in self.gunler for s in self.saatler}
        secili = self.tree_sinif.selection()
        if secili:
            idx = self.tree_sinif.index(secili[0])
            self.siniflar[idx] = {"ad": ad, "tip": s_tip, "musaitlik": musaitlik_durumu}
            self.tree_sinif.item(secili[0], values=(ad, s_tip))
        else:
            if any(s["ad"] == ad for s in self.siniflar): return messagebox.showwarning("Uyarı", "Bu mekan zaten var.")
            self.siniflar.append({"ad": ad, "tip": s_tip, "musaitlik": musaitlik_durumu})
            self.tree_sinif.insert("", tk.END, values=(ad, s_tip))
        self.veri_kaydet()

    def sinif_sil(self):
        secili = self.tree_sinif.selection()
        if not secili: return
        idx = self.tree_sinif.index(secili[0])
        del self.siniflar[idx]
        self.tree_sinif.delete(secili[0])
        self.veri_kaydet()
        self.form_sifirla_sinif()

    def ders_toplu_durum(self, durum):
        aktif_agac = self.aktif_tabloyu_getir()
        secililer = aktif_agac.selection()
        if not secililer: return messagebox.showwarning("Uyarı", "Lütfen tablodan en az bir ders seçin.")
        for item in secililer:
            idx = int(aktif_agac.item(item, "tags")[0])
            self.dersler[idx]["aktif_mi"] = durum
        self.veri_kaydet()
        self.ders_tablolarini_yenile()

    def ders_ekle(self):
        ad = self.entry_ders.get().strip()
        secili_hocalar_idx = self.listbox_ogr.curselection()
        if not secili_hocalar_idx: return messagebox.showwarning("Uyarı", "En az bir Öğretim Elemanı seçin.")
        secilen_hocalar = [self.listbox_ogr.get(i) for i in secili_hocalar_idx]
        
        yeni_ders = {
            "id": ad, "hocalar": secilen_hocalar, 
            "sinif_yil": self.combo_sinif_yili.get(), "mufredat": self.combo_mufredat.get(), 
            "donem": self.combo_donem_kayit.get(), "derslik_gerekmez": self.derslik_gerekmez_var.get(), 
            "istenen_mekan": self.combo_istenen_mekan.get(), "aktif_mi": self.aktif_mi_var.get(),
            "sabit_gun": self.combo_sabit_gun.get(), "sabit_saat": self.combo_sabit_saat.get()
        }
        
        try:
            yeni_ders["sure"] = int(self.entry_sure.get().strip())
            if not ad: raise ValueError
            
            aktif_agac = self.aktif_tabloyu_getir()
            secili = aktif_agac.selection()
            
            if len(secili) == 1:
                idx = int(aktif_agac.item(secili[0], "tags")[0])
                self.dersler[idx] = yeni_ders
            else:
                self.dersler.append(yeni_ders)
                
            self.veri_kaydet()
            self.ders_tablolarini_yenile()
        except:
            messagebox.showerror("Hata", "Eksik/hatalı giriş. (Süreye rakam yazınız).")

    def ders_sil(self):
        aktif_agac = self.aktif_tabloyu_getir()
        secililer = aktif_agac.selection()
        if not secililer: return
        indeksler = sorted([int(aktif_agac.item(item, "tags")[0]) for item in secililer], reverse=True)
        for idx in indeksler:
            del self.dersler[idx]
        self.veri_kaydet()
        self.form_sifirla_ders()
        self.ders_tablolarini_yenile()

    def program_uret(self):
        secilen_donem = self.combo_donem_secim.get()
        aktif_dersler = [d for d in self.dersler if d["donem"] == secilen_donem and d.get("aktif_mi", True)]
        
        if not aktif_dersler: return messagebox.showwarning("Uyarı", f"'{secilen_donem}' döneminde 'Aktif' olarak işaretlenmiş ders bulunamadı.")
        if not self.siniflar or not self.ogretim_elemanlari: return messagebox.showwarning("Uyarı", "Mekan veya öğretim elemanı eklemediniz.")

        onay_mesaji = f"Seçilen Dönem: {secilen_donem}\n\nAşağıdaki AKTİF dersler için program yapılacaktır:\n\n"
        for d in aktif_dersler: onay_mesaji += f"• {d['id']} ({', '.join(d['hocalar'])}) - {d['sure']} Saat\n"
        onay_mesaji += "\nOnaylıyor musunuz?"
        if not messagebox.askyesno("Onay", onay_mesaji): return

        for item in self.tree_sonuc.get_children(): self.tree_sonuc.delete(item)
        self.olusturulan_program = []
        model = cp_model.CpModel()
        gun_sayisi, gunluk_saat = len(self.gunler), len(self.saatler)
        
        ders_baslangic = {}
        ogr_araliklari = {o["ad"]: [] for o in self.ogretim_elemanlari}
        sinif_araliklari = {s["ad"]: [] for s in self.siniflar}
        sinif_secimleri = {}
        
        ogrenci_araliklari = {"1. Sınıf": [], "2. Sınıf": [], "3. Sınıf": [], "4. Sınıf": [], "Seçmeli/Karma": []}
        
        for o in self.ogretim_elemanlari:
            for gun_idx, gun in enumerate(self.gunler):
                for saat_idx, saat in enumerate(self.saatler):
                    if not o["musaitlik"][(gun, saat)]:
                        zaman = gun_idx * gunluk_saat + saat_idx
                        ogr_araliklari[o["ad"]].append(model.NewIntervalVar(zaman, 1, zaman + 1, f'ko_{o["ad"]}_{zaman}'))

        for s in self.siniflar:
            for gun_idx, gun in enumerate(self.gunler):
                for saat_idx, saat in enumerate(self.saatler):
                    if not s["musaitlik"][(gun, saat)]:
                        zaman = gun_idx * gunluk_saat + saat_idx
                        sinif_araliklari[s["ad"]].append(model.NewIntervalVar(zaman, 1, zaman + 1, f'ks_{s["ad"]}_{zaman}'))

        for i, ders in enumerate(aktif_dersler):
            sure = ders["sure"]
            gecerli_saatler = []
            
            for hoca in ders["hocalar"]:
                if hoca not in ogr_araliklari: return messagebox.showerror("Hata", f"Öğretim Elemanı bulunamadı: {hoca}")

            sabit_gun = ders.get("sabit_gun", "(Esnek / Serbest)")
            sabit_saat = ders.get("sabit_saat", "(Esnek / Serbest)")

            if sabit_gun != "(Esnek / Serbest)" or sabit_saat != "(Esnek / Serbest)":
                try:
                    if sabit_gun != "(Esnek / Serbest)" and sabit_saat != "(Esnek / Serbest)":
                        g_idx = self.gunler.index(sabit_gun)
                        s_idx = self.saatler.index(sabit_saat)
                        sabit_zaman_index = g_idx * gunluk_saat + s_idx
                        if sure == 8:
                            if s_idx != 0: return messagebox.showerror("Hata", f"'{ders['id']}' 8 saatlik ders olduğu için saat 08:30 olmalıdır.")
                        else:
                            if s_idx + sure > gunluk_saat or (s_idx < self.ogle_arasi_siniri and (s_idx + sure) > self.ogle_arasi_siniri):
                                return messagebox.showerror("Hata", f"'{ders['id']}' sabit saati öğle arasına taşıyor!")
                        gecerli_saatler = [sabit_zaman_index]

                    elif sabit_gun != "(Esnek / Serbest)" and sabit_saat == "(Esnek / Serbest)":
                        g_idx = self.gunler.index(sabit_gun)
                        for saat_idx in range(gunluk_saat):
                            if sure == 8:
                                if saat_idx == 0: gecerli_saatler.append(g_idx * gunluk_saat + 0)
                            else:
                                if saat_idx + sure > gunluk_saat: continue
                                if saat_idx < self.ogle_arasi_siniri and (saat_idx + sure) > self.ogle_arasi_siniri: continue
                                gecerli_saatler.append(g_idx * gunluk_saat + saat_idx)

                    elif sabit_gun == "(Esnek / Serbest)" and sabit_saat != "(Esnek / Serbest)":
                        s_idx = self.saatler.index(sabit_saat)
                        if sure == 8 and s_idx != 0:
                            return messagebox.showerror("Hata", f"'{ders['id']}' 8 saatlik ders olduğu için saat 08:30 olmalıdır.")
                        else:
                            if s_idx + sure > gunluk_saat or (s_idx < self.ogle_arasi_siniri and (s_idx + sure) > self.ogle_arasi_siniri):
                                return messagebox.showerror("Hata", f"'{ders['id']}' sabit saati öğle arasına taşıyor!")
                        for g_idx in range(gun_sayisi):
                            gecerli_saatler.append(g_idx * gunluk_saat + s_idx)
                except:
                    gecerli_saatler = []

            if not gecerli_saatler:
                for gun_idx in range(gun_sayisi):
                    for saat_idx in range(gunluk_saat):
                        if sure == 8:
                            if saat_idx == 0: gecerli_saatler.append(gun_idx * gunluk_saat + 0)
                        else:
                            if saat_idx + sure > gunluk_saat: continue
                            if saat_idx < self.ogle_arasi_siniri and (saat_idx + sure) > self.ogle_arasi_siniri: continue
                            gecerli_saatler.append(gun_idx * gunluk_saat + saat_idx)

            if not gecerli_saatler: return messagebox.showerror("Hata", f"Ders sığmıyor veya kurallara uymuyor: {ders['id']}")

            baslangic = model.NewIntVarFromDomain(cp_model.Domain.FromValues(gecerli_saatler), f'bas_{i}')
            bitis = model.NewIntVar(0, gun_sayisi * gunluk_saat, f'bitis_{i}')
            aralik = model.NewIntervalVar(baslangic, sure, bitis, f'aralik_{i}')
            
            ders_baslangic[i] = baslangic
            for hoca in ders["hocalar"]: ogr_araliklari[hoca].append(aralik)

            sy = ders.get("sinif_yil", "1. Sınıf")
            if sy in ogrenci_araliklari:
                ogrenci_araliklari[sy].append(aralik)

            if not ders.get("derslik_gerekmez", False):
                ders_sinif_secimleri = []
                istenen_tip = ders.get("istenen_mekan", "Derslik (Genel)")
                
                for s in self.siniflar:
                    mekan_tipi = s.get("tip", "Derslik (Genel)")
                    if mekan_tipi != istenen_tip: continue

                    secildi_mi = model.NewBoolVar(f'sec_{i}_{s["ad"]}')
                    ders_sinif_secimleri.append(secildi_mi)
                    sinif_secimleri[(i, s["ad"])] = secildi_mi
                    sinif_araliklari[s["ad"]].append(model.NewOptionalIntervalVar(baslangic, sure, bitis, secildi_mi, f'ops_{i}_{s["ad"]}'))
                
                if not ders_sinif_secimleri:
                    return messagebox.showerror("Mekan Hatası", f"'{ders['id']}' için istenen türde ({istenen_tip}) kayıtlı boş mekan bulunamadı!")
                
                model.AddExactlyOne(ders_sinif_secimleri)

        for araliklar in ogr_araliklari.values(): model.AddNoOverlap(araliklar)
        for araliklar in sinif_araliklari.values(): model.AddNoOverlap(araliklar)
        for araliklar in ogrenci_araliklari.values(): 
            if araliklar: model.AddNoOverlap(araliklar)

        self.root.update()
        solver = cp_model.CpSolver()
        solver.parameters.max_time_in_seconds = 20.0 
        
        status = solver.Solve(model)

        if status == cp_model.UNKNOWN:
            messagebox.showerror("Zamanaşımı / Darboğaz", "Algoritma 20 saniye boyunca denedi ancak bu dar kısıtlamalarla bir program çıkaramadı.\n\nLütfen sabit gün/saat kilitlerini, mekan tiplerini ve öğretim elemanı müsaitliklerini kontrol edin.")
            return

        elif status == cp_model.OPTIMAL or status == cp_model.FEASIBLE:
            for i, ders in enumerate(aktif_dersler):
                bas = solver.Value(ders_baslangic[i])
                sure = ders["sure"]
                
                if ders.get("derslik_gerekmez", False):
                    atanan_sinif = "Belirtilmedi (Ofis/Dış)"
                else:
                    atanan_sinif = next((s["ad"] for s in self.siniflar if (i, s["ad"]) in sinif_secimleri and solver.Value(sinif_secimleri[(i, s["ad"])]) == 1), "Atanmadı")
                
                hoca_metni = ", ".join(ders["hocalar"])
                
                for saat_artisi in range(sure):
                    gecerli_zaman = bas + saat_artisi
                    gun = self.gunler[gecerli_zaman // gunluk_saat]
                    saat = self.saatler[gecerli_zaman % gunluk_saat]
                    
                    self.olusturulan_program.append({
                        "Sıra (Gizli)": gecerli_zaman, "Dönem": ders["donem"], "Müfredat": ders["mufredat"],
                        "Sınıf Seviyesi": ders["sinif_yil"], "Gün": gun, "Saat": saat, "Derslik": atanan_sinif,
                        "Ders": ders["id"], "Öğretim Elemanı": hoca_metni, "Süre": f"{ders['sure']} Saat"
                    })
            
            self.olusturulan_program.sort(key=lambda x: x["Sıra (Gizli)"])
            for v in self.olusturulan_program:
                self.tree_sonuc.insert("", tk.END, values=(v["Gün"], v["Saat"], v["Derslik"], v["Ders"], v["Sınıf Seviyesi"], v["Öğretim Elemanı"]))
            messagebox.showinfo("Başarılı", f"{secilen_donem} ders programı kısıtlara uygun olarak üretildi!")
        else:
            messagebox.showerror("Çözüm Bulunamadı", "Saatler ve ders yükü matematiksel olarak çakışıyor, program yerleştirilemedi. Sabit gün kilitlerinin veya istenen mekan tiplerinin çakışıp çakışmadığını kontrol edin.")

    def excel_aktar(self):
        if not self.olusturulan_program: 
            return messagebox.showwarning("Uyarı", "Önce programı üretmelisiniz.")
            
        dosya = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], title="Kaydet")
        if not dosya: return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, Border, Side

            wb = Workbook()
            ws = wb.active
            ws.title = "Ders Programı"
            
            sinif_seviyeleri = ["1. Sınıf", "2. Sınıf", "3. Sınıf", "4. Sınıf", "Seçmeli/Karma"]
            
            ws.merge_cells('A1:G1')
            baslik = ws['A1']
            baslik.value = f"KARAMANOĞLU MEHMETBEY ÜNİVERSİTESİ {self.bolum_adi} N.Ö"
            baslik.font = Font(bold=True, size=12)
            baslik.alignment = Alignment(horizontal="center", vertical="center")
            
            headers = ["GÜN", "SAAT"] + sinif_seviyeleri
            ws.append(headers)
            
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=2, column=col_num)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            matris = { (g, s): {lvl: "" for lvl in sinif_seviyeleri} for g in self.gunler for s in self.saatler }
            
            for p in self.olusturulan_program:
                g = p["Gün"]
                s = p["Saat"]
                lvl = p["Sınıf Seviyesi"]
                
                ders_ad = p["Ders"]
                if " " in ders_ad and ders_ad.split(" ")[0].isdigit():
                    ders_ad = " ".join(ders_ad.split(" ")[1:])
                    
                metin = f"{ders_ad}\n({p['Öğretim Elemanı']})\n{p['Derslik']}"
                
                if matris[(g, s)][lvl] == "":
                    matris[(g, s)][lvl] = metin
                else:
                    matris[(g, s)][lvl] += f"\n\n{metin}" 
                    
            border_style = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            
            row_idx = 3
            for g in self.gunler:
                start_row = row_idx
                for s in self.saatler:
                    row_data = [g, s] + [matris[(g, s)][lvl] for lvl in sinif_seviyeleri]
                    ws.append(row_data)
                    ws.row_dimensions[row_idx].height = 65
                    
                    for col_idx in range(1, 8):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        cell.border = border_style
                    row_idx += 1
                    
                ws.merge_cells(start_row=start_row, start_column=1, end_row=row_idx-1, end_column=1)
                ws.cell(row=start_row, column=1).alignment = Alignment(horizontal="center", vertical="center", textRotation=90)
                
            ws.column_dimensions['A'].width = 6
            ws.column_dimensions['B'].width = 15
            for c in ['C', 'D', 'E', 'F', 'G']:
                ws.column_dimensions[c].width = 25
                
            wb.save(dosya)
            messagebox.showinfo("Başarılı", f"Program, okulunuzun formatında Excel'e aktarıldı:\n{dosya}")

        except Exception as e:
            messagebox.showerror("Aktarma Hatası", f"Excel'e aktarılırken bir sorun oluştu:\n{str(e)}")

if __name__ == "__main__":
    root = tk.Tk()
    root.report_callback_exception = hata_yakalayici 
    app = DersProgramiApp(root)
    root.mainloop()